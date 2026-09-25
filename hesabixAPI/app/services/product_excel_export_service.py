"""
ساخت فایل Excel خروجی لیست محصولات (sync و background job).
"""

from __future__ import annotations

import datetime
import io
import logging
import re
from types import SimpleNamespace
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import Workbook
from openpyxl.cell.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.orm import Session

from adapters.db.models.business import Business
from adapters.db.models.currency import Currency
from adapters.db.models.price_list import PriceItem, PriceList
from app.core.i18n import negotiate_locale
from app.core.responses import format_datetime_fields

logger = logging.getLogger(__name__)

PRODUCT_EXPORT_INVENTORY_KEYS = frozenset({
	"inventory_stock_warehouse",
	"inventory_stock_accounting",
	"inventory_stock_physical",
	"inventory_stock_financial",
	"warehouse_recharge",
})

MAX_EXPORT_RECORDS = 10000


def body_truthy(value: Any) -> bool:
	if isinstance(value, bool):
		return value
	if value is None:
		return False
	if isinstance(value, (int, float)):
		return value != 0
	if isinstance(value, str):
		return value.strip().lower() in {"1", "true", "yes", "on"}
	return bool(value)


def export_columns_need_inventory(body: dict) -> bool:
	"""اگر ستون‌های خروجی به موجودی وابسته‌اند، محاسبهٔ موجودی لازم است."""
	export_columns = body.get("export_columns")
	if not isinstance(export_columns, list):
		return False
	for col in export_columns:
		if isinstance(col, dict) and col.get("key") in PRODUCT_EXPORT_INVENTORY_KEYS:
			return True
	return False


def should_include_inventory_for_product_export(body: dict) -> bool:
	"""
	وقتی export_columns ارسال شده، فقط بر اساس ستون‌ها تصمیم بگیر
	(تا include_inventory سراسری UI باعث محاسبهٔ سنگین نشود).
	"""
	export_columns = body.get("export_columns")
	if isinstance(export_columns, list) and len(export_columns) > 0:
		return export_columns_need_inventory(body)
	if body_truthy(body.get("include_inventory")):
		return True
	return False


def enrich_product_export_items(items: List[dict], *, is_fa: bool = True) -> None:
	"""پر کردن فیلدهای محاسباتی مخصوص خروجی و نرمال‌سازی نمایش."""
	yes_label = "بله" if is_fa else "Yes"
	no_label = "خیر" if is_fa else "No"
	dash = "-"

	for it in items:
		if not isinstance(it, dict):
			continue

		track = bool(it.get("track_inventory"))

		if "inventory_stock_warehouse" not in it and "inventory_stock_physical" in it:
			it["inventory_stock_warehouse"] = it.get("inventory_stock_physical")
		if "inventory_stock_accounting" not in it and "inventory_stock_financial" in it:
			it["inventory_stock_accounting"] = it.get("inventory_stock_financial")

		wh_stock = it.get("inventory_stock_warehouse")
		acc_stock = it.get("inventory_stock_accounting")
		if not track:
			it["inventory_stock_warehouse"] = dash
			it["inventory_stock_accounting"] = dash
			it["warehouse_recharge"] = dash
		else:
			if wh_stock is None:
				it["inventory_stock_warehouse"] = dash
			if acc_stock is None:
				it["inventory_stock_accounting"] = dash

			reorder_point = it.get("reorder_point")
			stock_for_reorder = it.get("inventory_stock_accounting")
			if stock_for_reorder == dash or reorder_point is None:
				it["warehouse_recharge"] = no_label
			else:
				try:
					stock_num = float(stock_for_reorder)
					reorder_num = float(reorder_point)
					it["warehouse_recharge"] = yes_label if stock_num < reorder_num else no_label
				except (TypeError, ValueError):
					it["warehouse_recharge"] = no_label

		if "track_inventory" in it:
			it["track_inventory"] = yes_label if track else no_label

		wh_name = it.get("default_warehouse_name")
		wh_code = it.get("default_warehouse_code")
		if wh_name and wh_code and not it.get("default_warehouse_display"):
			it["default_warehouse_name"] = f"{wh_code} - {wh_name}"

		if it.get("image") in (None, ""):
			it["image"] = it.get("thumbnail_url") or it.get("image_url") or ""


def _fake_request(*, calendar_type: str, accept_language: str) -> Any:
	return SimpleNamespace(
		state=SimpleNamespace(calendar_type=calendar_type or "gregorian"),
		headers={"Accept-Language": accept_language or "en"},
	)


def _slugify(text: str) -> str:
	return re.sub(r"[^A-Za-z0-9_-]+", "_", text).strip("_")


def _format_price_list_value(rows) -> str:
	if not rows:
		return ""
	parts = []
	for price, currency_code, tier_name, min_qty in rows:
		tier_label = (tier_name or "").strip() or "base"
		cur_label = (currency_code or "").strip() or "-"
		qty_text = ""
		try:
			if min_qty is not None and float(min_qty) > 0:
				qty_text = f" (>= {min_qty})"
		except Exception:
			qty_text = ""
		parts.append(f"{tier_label}{qty_text} [{cur_label}]: {price}")
	return " | ".join(parts)


def _apply_selected_filters(items: List[dict], body: dict) -> List[dict]:
	selected_only = bool(body.get("selected_only", False))
	selected_row_keys = body.get("selected_row_keys")
	selected_indices = body.get("selected_indices")

	if selected_only and isinstance(selected_row_keys, list):
		try:
			wanted_ids = set()
			for key in selected_row_keys:
				if isinstance(key, dict) and key.get("id") is not None:
					wanted_ids.add(int(key.get("id")))
			if wanted_ids:
				filtered = []
				for it in items:
					try:
						if int(it.get("id")) in wanted_ids:
							filtered.append(it)
					except Exception:
						continue
				items = filtered
		except Exception:
			pass

	if selected_only and selected_indices is not None and (
		not isinstance(selected_row_keys, list) or not selected_row_keys or not items
	):
		indices = None
		if isinstance(selected_indices, str):
			try:
				import json as _json
				indices = _json.loads(selected_indices)
			except Exception:
				indices = None
		elif isinstance(selected_indices, list):
			indices = selected_indices
		if isinstance(indices, list):
			items = [items[i] for i in indices if isinstance(i, int) and 0 <= i < len(items)]
	return items


def _attach_price_list_columns(
	db: Session,
	business_id: int,
	items: List[dict],
	body: dict,
) -> List[Tuple[str, str]]:
	report_mode = str(body.get("report_mode") or "base_plus_price_lists").strip().lower()
	if report_mode not in {"base_only", "price_lists_only", "base_plus_price_lists"}:
		report_mode = "base_plus_price_lists"

	raw_price_list_ids = body.get("price_list_ids") or []
	price_list_ids: List[int] = []
	if isinstance(raw_price_list_ids, list):
		seen = set()
		for v in raw_price_list_ids:
			try:
				pid = int(v)
			except Exception:
				continue
			if pid not in seen:
				seen.add(pid)
				price_list_ids.append(pid)

	dynamic_price_columns: List[Tuple[str, str]] = []
	if report_mode not in {"price_lists_only", "base_plus_price_lists"} or not price_list_ids or not items:
		return dynamic_price_columns

	product_ids = []
	for it in items:
		try:
			product_ids.append(int(it.get("id")))
		except Exception:
			continue
	if not product_ids:
		return dynamic_price_columns

	pl_rows = db.query(PriceList.id, PriceList.name).filter(
		PriceList.business_id == business_id,
		PriceList.id.in_(price_list_ids),
	).all()
	price_list_name_by_id = {int(pid): (name or f"#{pid}") for pid, name in pl_rows}

	pi_rows = (
		db.query(
			PriceItem.product_id,
			PriceItem.price_list_id,
			PriceItem.price,
			Currency.code,
			PriceItem.tier_name,
			PriceItem.min_qty,
		)
		.join(PriceList, PriceList.id == PriceItem.price_list_id)
		.outerjoin(Currency, Currency.id == PriceItem.currency_id)
		.filter(
			PriceList.business_id == business_id,
			PriceItem.product_id.in_(product_ids),
			PriceItem.price_list_id.in_(price_list_ids),
		)
		.order_by(PriceItem.price_list_id.asc(), PriceItem.min_qty.asc(), PriceItem.tier_name.asc())
		.all()
	)

	grouped: Dict[Tuple[int, int], list] = {}
	for product_id, price_list_id, price, currency_code, tier_name, min_qty in pi_rows:
		key = (int(product_id), int(price_list_id))
		grouped.setdefault(key, []).append((price, currency_code, tier_name, min_qty))

	for price_list_id in price_list_ids:
		if price_list_id not in price_list_name_by_id:
			continue
		col_key = f"price_list_{price_list_id}"
		col_label = f"لیست قیمت: {price_list_name_by_id[price_list_id]}"
		dynamic_price_columns.append((col_key, col_label))
		for it in items:
			try:
				pid = int(it.get("id"))
			except Exception:
				continue
			it[col_key] = _format_price_list_value(grouped.get((pid, price_list_id), []))

	return dynamic_price_columns


def _resolve_headers_keys(body: dict, dynamic_price_columns: List[Tuple[str, str]]) -> Tuple[List[str], List[str]]:
	report_mode = str(body.get("report_mode") or "base_plus_price_lists").strip().lower()
	if report_mode not in {"base_only", "price_lists_only", "base_plus_price_lists"}:
		report_mode = "base_plus_price_lists"

	columns_profile = str(body.get("columns_profile") or "").strip().lower()
	force_price_report_columns = columns_profile in {"price_report", "product_price_report"}

	export_columns = body.get("export_columns")
	if export_columns and isinstance(export_columns, list) and not force_price_report_columns:
		headers = [col.get("label") or col.get("key") for col in export_columns]
		keys = [col.get("key") for col in export_columns]
		return headers, keys

	default_cols = [
		("code", "کد"),
		("name", "نام"),
		("category_name", "دسته"),
	]
	if report_mode in {"base_only", "base_plus_price_lists"}:
		default_cols.extend([
			("base_sales_price", "قیمت فروش پایه"),
			("base_purchase_price", "قیمت خرید پایه"),
		])
	if dynamic_price_columns:
		default_cols.extend(dynamic_price_columns)
	if not force_price_report_columns:
		default_cols.extend([
			("main_unit", "واحد اصلی"),
			("secondary_unit", "واحد فرعی"),
			("track_inventory", "کنترل موجودی"),
			("created_at_formatted", "ایجاد"),
		])
	keys = [k for k, _ in default_cols]
	headers = [v for _, v in default_cols]
	return headers, keys


def load_product_export_items(
	db: Session,
	business_id: int,
	body: dict,
	*,
	calendar_type: str = "jalali",
	accept_language: str = "fa",
) -> List[dict]:
	selected_only = bool(body.get("selected_only", False))
	selected_row_keys = body.get("selected_row_keys")
	has_selected_row_keys = (
		selected_only
		and isinstance(selected_row_keys, list)
		and any(isinstance(k, dict) and k.get("id") is not None for k in selected_row_keys)
	)

	if selected_only and not has_selected_row_keys:
		take = max(1, min(int(body.get("take", 1000)), MAX_EXPORT_RECORDS))
		skip = max(0, int(body.get("skip", 0)))
	else:
		take = MAX_EXPORT_RECORDS
		skip = 0

	include_inventory = should_include_inventory_for_product_export(body)
	query_dict = {
		"take": take,
		"skip": skip,
		"sort_by": body.get("sort_by"),
		"sort_desc": bool(body.get("sort_desc", False)),
		"sort": body.get("sort") if isinstance(body.get("sort"), list) else None,
		"search": body.get("search"),
		"search_fields": body.get("search_fields") or body.get("searchFields"),
		"filters": body.get("filters"),
		"category_ids": body.get("category_ids") or body.get("categoryIds"),
		"include_inventory": include_inventory,
		"inventory_as_of_date": body.get("inventory_as_of_date") or body.get("inventoryAsOfDate"),
	}
	from app.services.product_service import list_products

	result = list_products(db, business_id, query_dict)
	items = result.get("items", []) if isinstance(result, dict) else []

	fake_req = _fake_request(calendar_type=calendar_type, accept_language=accept_language)
	items = [format_datetime_fields(item, fake_req, business_id=business_id) for item in items]
	items = _apply_selected_filters(items, body)

	max_items = body.get("limit")
	try:
		max_items = int(max_items) if max_items is not None else None
	except Exception:
		max_items = None
	if max_items is not None:
		max_items = max(1, min(max_items, MAX_EXPORT_RECORDS))
		items = items[:max_items]

	return items


def build_products_excel_export(
	db: Session,
	business_id: int,
	body: dict,
	*,
	calendar_type: Optional[str] = None,
	accept_language: Optional[str] = None,
) -> Tuple[bytes, str, int]:
	"""
	ساخت فایل Excel محصولات.

	Returns:
		(xlsx_bytes, filename, record_count)
	"""
	accept_language = accept_language or "fa"
	locale = negotiate_locale(accept_language)
	if not calendar_type:
		calendar_type = "jalali" if locale == "fa" else "gregorian"

	items = load_product_export_items(
		db,
		business_id,
		body,
		calendar_type=calendar_type,
		accept_language=accept_language,
	)
	dynamic_price_columns = _attach_price_list_columns(db, business_id, items, body)
	headers, keys = _resolve_headers_keys(body, dynamic_price_columns)

	enrich_product_export_items(items, is_fa=(locale == "fa"))

	# write_only: حافظه کمتر برای کاتالوگ‌های بزرگ؛ استایل فقط برای هدر
	wb = Workbook(write_only=True)
	ws = wb.create_sheet(title="Products")
	if locale == "fa":
		try:
			ws.sheet_view.rightToLeft = True
		except Exception:
			pass

	header_font = Font(bold=True)
	header_fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")
	header_cells = []
	for h in headers:
		cell = WriteOnlyCell(ws, value=h)
		cell.font = header_font
		cell.fill = header_fill
		cell.alignment = Alignment(horizontal="center")
		header_cells.append(cell)
	ws.append(header_cells)

	for it in items:
		ws.append([it.get(k) for k in keys])

	output = io.BytesIO()
	wb.save(output)
	data = output.getvalue()

	biz_name = ""
	try:
		b = db.query(Business).filter(Business.id == business_id).first()
		if b is not None:
			biz_name = b.name or ""
	except Exception:
		biz_name = ""

	base = "products"
	if biz_name:
		base += f"_{_slugify(biz_name)}"
	if bool(body.get("selected_only", False)):
		base += "_selected"
	filename = f"{base}_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
	return data, filename, len(items)


def upload_export_bytes(
	db: Session,
	*,
	business_id: int,
	user_id: int,
	data: bytes,
	filename: str,
) -> Dict[str, Any]:
	"""آپلود فایل خروجی به ذخیره‌ساز کسب‌وکار (برای دانلود بعدی)."""
	import anyio
	from fastapi import UploadFile
	from app.services.file_storage_service import FileStorageService

	faux_upload = UploadFile(filename=filename, file=io.BytesIO(data))
	storage = FileStorageService(db)

	async def _upload():
		return await storage.upload_file(
			faux_upload,
			user_id=user_id,
			module_context="product_export",
			developer_data={"business_id": business_id, "kind": "products_excel"},
			is_temporary=True,
			expires_in_days=7,
			business_id=business_id,
			check_storage_limit=False,
		)

	saved = anyio.run(_upload)
	return saved or {}
