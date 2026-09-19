"""خروجی Excel/PDF عمومی برای گزارش‌های لیستی."""
from __future__ import annotations

import io
import re
from dataclasses import dataclass
from html import escape
from typing import Any, Callable, Dict, List, Optional, Sequence, Tuple

from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from weasyprint import CSS, HTML
from weasyprint.text.fonts import FontConfiguration

from adapters.db.models.business import Business


@dataclass(frozen=True)
class ListReportColumn:
    key: str
    label_fa: str
    label_en: str
    kind: str = "text"  # text | number | date


def _slugify(text: str) -> str:
    return re.sub(r"[^A-Za-z0-9_-]+", "_", str(text)).strip("_")


def _cell_value(item: Dict[str, Any], col: ListReportColumn) -> Any:
    val = item.get(col.key)
    if val is None:
        return ""
    if col.kind == "number":
        try:
            f = float(val)
            return int(f) if f == int(f) else f
        except (TypeError, ValueError):
            return val
    if col.kind == "date":
        s = str(val)
        return s.split("T")[0].split(" ")[0]
    return val


def _format_display(value: Any, is_fa: bool, kind: str) -> str:
    if value is None or value == "":
        return ""
    if kind == "number":
        try:
            f = float(value)
            s = f"{int(f):,}" if f == int(f) else f"{f:,.2f}"
            return s.replace(",", "٬") if is_fa else s
        except (TypeError, ValueError):
            return str(value)
    return str(value)


def build_list_excel_bytes(
    items: Sequence[Dict[str, Any]],
    columns: Sequence[ListReportColumn],
    *,
    is_fa: bool,
    title: str,
    summary: Optional[Dict[str, Any]] = None,
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "گزارش" if is_fa else "Report"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    thin = Side(border_style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    if title:
        ws.append([title])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(columns), 1))
        ws.cell(row=1, column=1).font = Font(bold=True, size=13)

    headers = [c.label_fa if is_fa else c.label_en for c in columns]
    ws.append(headers)
    header_row = ws.max_row
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=header_row, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = center

    for item in items:
        ws.append([_cell_value(item, c) for c in columns])

    for r in range(header_row + 1, ws.max_row + 1):
        for c in range(1, len(columns) + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = border
            cell.alignment = center

    if summary:
        ws.append([])
        ws.append(["خلاصه" if is_fa else "Summary"])
        for k, v in summary.items():
            ws.append([k, v])

    for column in ws.columns:
        letter = column[0].column_letter
        max_len = 0
        for cell in column:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max_len + 2, 50)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def build_list_pdf_bytes(
    items: Sequence[Dict[str, Any]],
    columns: Sequence[ListReportColumn],
    *,
    is_fa: bool,
    title: str,
    business_name: str,
    generated_at: str,
    summary: Optional[Dict[str, Any]] = None,
    load_farsi_font_data_uris: Optional[Callable[[], Tuple[Optional[str], Optional[str]]]] = None,
) -> bytes:
    def esc(v: Any) -> str:
        return escape("" if v is None else str(v))

    headers = [c.label_fa if is_fa else c.label_en for c in columns]
    rows_html = []
    for item in items:
        cells = "".join(
            f"<td class='{'num' if c.kind == 'number' else 'txt'}'>"
            f"{esc(_format_display(_cell_value(item, c), is_fa, c.kind))}</td>"
            for c in columns
        )
        rows_html.append(f"<tr>{cells}</tr>")

    summary_html = ""
    if summary:
        parts = "".join(
            f"<div class='sum-item'><span>{esc(k)}</span><strong>{esc(_format_display(v, is_fa, 'number'))}</strong></div>"
            for k, v in summary.items()
        )
        summary_html = f"<div class='summary'>{parts}</div>"

    html_lang = "fa" if is_fa else "en"
    html_dir = "rtl" if is_fa else "ltr"
    preferred = "YekanBakhFaNum, Vazirmatn, Tahoma, Arial, sans-serif" if is_fa else "Arial, sans-serif"

    table_html = f"""
    <html lang="{html_lang}" dir="{html_dir}">
      <head><meta charset="utf-8"></head>
      <body>
        <div class="header">
          <div class="title">{esc(title)}</div>
          <div class="meta">
            <div>{esc(business_name)}</div>
            <div>{esc(generated_at)}</div>
          </div>
        </div>
        {summary_html}
        <table>
          <thead><tr>{''.join(f'<th>{esc(h)}</th>' for h in headers)}</tr></thead>
          <tbody>{''.join(rows_html)}</tbody>
        </table>
      </body>
    </html>
    """

    font_face = ""
    if is_fa and load_farsi_font_data_uris:
        try:
            reg, bold = load_farsi_font_data_uris()
            if reg:
                font_face += f'@font-face {{ font-family: YekanBakhFaNum; src: url("{reg}") format("truetype"); font-weight:400; }}\n'
            if bold:
                font_face += f'@font-face {{ font-family: YekanBakhFaNum; src: url("{bold}") format("truetype"); font-weight:700; }}\n'
        except Exception:
            pass

    css = CSS(string=font_face + f"""
      @page {{ size: A4 landscape; margin: 10mm; }}
      body {{ font-family: {preferred}; font-size: 10px; color: #222; }}
      .header {{ display:flex; justify-content:space-between; margin-bottom:10px; border-bottom:2px solid #444; padding-bottom:6px; }}
      .title {{ font-size:14px; font-weight:700; }}
      .meta {{ font-size:10px; color:#555; text-align: {'right' if is_fa else 'left'}; }}
      .summary {{ display:flex; flex-wrap:wrap; gap:8px; margin-bottom:10px; }}
      .sum-item {{ background:#f8fafc; border:1px solid #e2e8f0; padding:4px 8px; border-radius:6px; }}
      table {{ width:100%; border-collapse:collapse; }}
      th {{ background:#f0f3f7; border:1px solid #c7cdd6; padding:5px; }}
      td {{ border:1px solid #d7dde6; padding:4px; }}
      td.num {{ text-align:center; }}
    """)
    return HTML(string=table_html).write_pdf(stylesheets=[css], font_config=FontConfiguration())


def list_excel_response(
    items: Sequence[Dict[str, Any]],
    columns: Sequence[ListReportColumn],
    *,
    db,
    business_id: int,
    filename_prefix: str,
    is_fa: bool,
    title: str,
    summary: Optional[Dict[str, Any]] = None,
    export_filename_timestamp: Callable[[int], str],
) -> Response:
    biz_name = ""
    try:
        b = db.query(Business).filter(Business.id == business_id).first()
        if b:
            biz_name = b.name or ""
    except Exception:
        pass
    content = build_list_excel_bytes(items, columns, is_fa=is_fa, title=title, summary=summary)
    filename = f"{filename_prefix}_{_slugify(biz_name) or business_id}_{export_filename_timestamp(business_id)}.xlsx"
    return Response(
        content=content,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename={filename}",
            "Content-Length": str(len(content)),
            "Access-Control-Expose-Headers": "Content-Disposition",
        },
    )


def list_pdf_response(
    items: Sequence[Dict[str, Any]],
    columns: Sequence[ListReportColumn],
    *,
    db,
    business_id: int,
    filename_prefix: str,
    is_fa: bool,
    title: str,
    generated_at: str,
    summary: Optional[Dict[str, Any]] = None,
    export_filename_timestamp: Callable[[int], str],
    load_farsi_font_data_uris: Callable[[], Tuple[Optional[str], Optional[str]]],
) -> Response:
    biz_name = ""
    try:
        b = db.query(Business).filter(Business.id == business_id).first()
        if b:
            biz_name = b.name or ""
    except Exception:
        pass
    pdf = build_list_pdf_bytes(
        items,
        columns,
        is_fa=is_fa,
        title=title,
        business_name=biz_name,
        generated_at=generated_at,
        summary=summary,
        load_farsi_font_data_uris=load_farsi_font_data_uris,
    )
    filename = f"{filename_prefix}_{_slugify(biz_name) or business_id}_{export_filename_timestamp(business_id)}.pdf"
    return Response(
        content=pdf,
        media_type="application/pdf",
        headers={
            "Content-Disposition": f"attachment; filename={filename}",
            "Content-Length": str(len(pdf)),
            "Access-Control-Expose-Headers": "Content-Disposition",
        },
    )
