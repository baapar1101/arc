"""خروجی Excel/PDF گزارش ترازنامه."""
from __future__ import annotations

import io
import re
from html import escape
from typing import Any, Callable, Dict, List, Optional, Tuple

from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy.orm import Session
from weasyprint import CSS, HTML
from weasyprint.text.fonts import FontConfiguration

from adapters.db.models.business import Business
from adapters.db.models.currency import Currency
from adapters.db.models.fiscal_year import FiscalYear


def _slugify(text: str) -> str:
    return re.sub(r"[^A-Za-z0-9_-]+", "_", str(text)).strip("_")


def _format_number_for_export(raw_value: Any, is_fa: bool) -> str:
    if raw_value is None:
        return "0"
    try:
        f = float(raw_value)
        if f == int(f):
            s = f"{int(f):,}"
        else:
            s = f"{f:,.2f}"
    except Exception:
        s = str(raw_value)
    if is_fa:
        s = s.replace(",", "٬")
    return s


def build_balance_sheet_excel_bytes(result: Dict[str, Any], is_fa: bool) -> bytes:
    summary = result.get("summary") or {}
    statement_lines = result.get("statement_lines") or []
    comparison = result.get("comparison") or {}
    has_compare = bool(comparison)

    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2E5C8A", end_color="2E5C8A", fill_type="solid")
    section_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
    subtotal_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    highlight_fill = PatternFill(start_color="E8F4FD", end_color="E8F4FD", fill_type="solid")
    balanced_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    unbalanced_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    thin = Side(border_style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws_summary = wb.active
    ws_summary.title = "خلاصه" if is_fa else "Summary"
    if has_compare:
        ws_summary.append(
            ["عنوان", "دوره جاری", "دوره قبل", "تغییر"]
            if is_fa
            else ["Label", "Current", "Prior", "Variance"]
        )
    else:
        ws_summary.append(["عنوان", "مبلغ"] if is_fa else ["Label", "Amount"])
    header_cols = 4 if has_compare else 2
    for col_idx in range(1, header_cols + 1):
        cell = ws_summary.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = center

    summary_keys = [
        ("total_assets", "جمع دارایی‌ها", "Total Assets"),
        ("total_liabilities", "جمع بدهی‌ها", "Total Liabilities"),
        ("current_period_profit", "سود (زیان) دوره جاری", "Current Period Profit"),
        ("total_equity", "جمع حقوق صاحبان سهام", "Total Equity"),
        ("total_liabilities_and_equity", "جمع بدهی و حقوق صاحبان سهام", "Total Liab. & Equity"),
        ("balance_difference", "اختلاف معادله", "Equation Difference"),
    ]
    for key, fa, en in summary_keys:
        label = fa if is_fa else en
        if has_compare and key in comparison:
            comp = comparison[key]
            ws_summary.append([label, comp.get("current"), comp.get("prior"), comp.get("variance")])
        else:
            ws_summary.append([label, summary.get(key, 0)])

    for r in range(2, ws_summary.max_row + 1):
        for c in range(1, header_cols + 1):
            cell = ws_summary.cell(row=r, column=c)
            cell.border = border
            cell.alignment = center

    ws_stmt = wb.create_sheet("ترازنامه" if is_fa else "Balance Sheet")
    if has_compare:
        stmt_headers = (
            ["شرح", "کد حساب", "نام حساب", "دوره جاری", "دوره قبل", "تغییر"]
            if is_fa
            else ["Description", "Code", "Account", "Current", "Prior", "Variance"]
        )
    else:
        stmt_headers = (
            ["شرح", "کد حساب", "نام حساب", "مبلغ"]
            if is_fa
            else ["Description", "Code", "Account", "Amount"]
        )
    ws_stmt.append(stmt_headers)
    for col_idx in range(1, len(stmt_headers) + 1):
        cell = ws_stmt.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = center

    for line in statement_lines:
        line_type = line.get("type")
        if line_type == "section_header":
            label = line.get("label_fa" if is_fa else "label_en", "")
            ws_stmt.append([label] + [""] * (len(stmt_headers) - 1))
            row = ws_stmt.max_row
            for c in range(1, len(stmt_headers) + 1):
                cell = ws_stmt.cell(row=row, column=c)
                cell.fill = section_fill
                cell.font = Font(bold=True)
        elif line_type == "account":
            if has_compare:
                ws_stmt.append([
                    "",
                    line.get("account_code", ""),
                    line.get("account_name", ""),
                    line.get("amount", 0),
                    line.get("prior_amount", 0),
                    line.get("variance", 0),
                ])
            else:
                ws_stmt.append(["", line.get("account_code", ""), line.get("account_name", ""), line.get("amount", 0)])
        elif line_type in ("subtotal", "grand_total"):
            label = line.get("label_fa" if is_fa else "label_en", "")
            if has_compare:
                ws_stmt.append([
                    label, "", "",
                    line.get("amount", 0),
                    line.get("prior_amount", 0),
                    line.get("variance", 0),
                ])
            else:
                ws_stmt.append([label, "", "", line.get("amount", 0)])
            row = ws_stmt.max_row
            fill = highlight_fill if line.get("highlight") else subtotal_fill
            for c in range(1, len(stmt_headers) + 1):
                cell = ws_stmt.cell(row=row, column=c)
                cell.fill = fill
                cell.font = Font(bold=True)
        elif line_type == "equation_check":
            label = line.get("label_fa" if is_fa else "label_en", "")
            ws_stmt.append([label, "", "", line.get("amount", 0)])
            row = ws_stmt.max_row
            fill = balanced_fill if line.get("equation_balanced") else unbalanced_fill
            for c in range(1, len(stmt_headers) + 1):
                cell = ws_stmt.cell(row=row, column=c)
                cell.fill = fill
                cell.font = Font(bold=True)

    for r in range(2, ws_stmt.max_row + 1):
        for c in range(1, len(stmt_headers) + 1):
            ws_stmt.cell(row=r, column=c).border = border
            ws_stmt.cell(row=r, column=c).alignment = center

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def _build_filters_html(
    db: Session,
    business_id: int,
    is_fa: bool,
    fiscal_year_id: Optional[int],
    date_from: Optional[str],
    date_to: Optional[str],
    currency_id: Optional[int],
    project_id: Optional[int],
    account_level: Optional[int] = None,
) -> str:
    filters: List[Tuple[str, str]] = []
    if fiscal_year_id:
        fy_title = str(fiscal_year_id)
        try:
            fy_obj = (
                db.query(FiscalYear)
                .filter(FiscalYear.id == int(fiscal_year_id), FiscalYear.business_id == int(business_id))
                .first()
            )
            if fy_obj:
                fy_title = fy_obj.title or fy_title
        except Exception:
            pass
        filters.append(("سال مالی" if is_fa else "Fiscal Year", fy_title))
    if date_to:
        filters.append(("تاریخ مبنا" if is_fa else "As of", date_to))
    if date_from:
        filters.append(("از تاریخ" if is_fa else "From", date_from))
    if currency_id:
        cur_label = str(currency_id)
        try:
            cur = db.query(Currency).filter(Currency.id == int(currency_id)).first()
            if cur is not None:
                cur_label = (cur.title or cur.name or "") if is_fa else (cur.code or cur.name or cur.title or "")
        except Exception:
            pass
        filters.append(("ارز" if is_fa else "Currency", cur_label))
    if project_id:
        filters.append(("پروژه" if is_fa else "Project", str(project_id)))
    if account_level:
        level_labels = {1: "گروه", 2: "کل", 3: "معین", 4: "تفصیل"}
        filters.append(("سطح نمایش" if is_fa else "Level", level_labels.get(int(account_level), str(account_level))))

    if not filters:
        return ""
    parts = [
        f"<span class='filter-item'><strong>{escape(k)}:</strong> {escape(v)}</span>"
        for k, v in filters
        if v
    ]
    return f"<div class='filters'>{''.join(parts)}</div>" if parts else ""


def build_balance_sheet_pdf_bytes(
    result: Dict[str, Any],
    *,
    is_fa: bool,
    business_name: str,
    report_title_fa: str,
    report_title_en: str,
    generated_at: str,
    filters_html: str,
    load_farsi_font_data_uris: Callable[[], Tuple[Optional[str], Optional[str]]],
) -> bytes:
    summary = result.get("summary") or {}
    statement_lines = result.get("statement_lines") or []
    comparison = result.get("comparison") or {}
    has_compare = bool(comparison)
    equation_ok = bool(summary.get("equation_balanced"))

    def esc(v: Any) -> str:
        return escape("" if v is None else str(v))

    def fmt(raw_value: Any) -> str:
        return esc(_format_number_for_export(raw_value, is_fa))

    title_text = report_title_fa if is_fa else report_title_en
    html_lang = "fa" if is_fa else "en"
    html_dir = "rtl" if is_fa else "ltr"

    summary_cards = [
        ("جمع دارایی‌ها" if is_fa else "Total Assets", summary.get("total_assets", 0), "assets"),
        ("جمع بدهی‌ها" if is_fa else "Total Liabilities", summary.get("total_liabilities", 0), "liab"),
        ("حقوق صاحبان سهام" if is_fa else "Equity", summary.get("total_equity", 0), "equity"),
        (
            "وضعیت معادله" if is_fa else "Equation",
            "متوازن" if equation_ok else f"اختلاف: {summary.get('balance_difference', 0)}",
            "balanced" if equation_ok else "unbalanced",
        ),
    ]
    summary_html = "<div class='summary'>" + "".join(
        f"<div class='summary-card {css}'><div class='label'>{esc(lbl)}</div>"
        f"<div class='value'>{fmt(val) if isinstance(val, (int, float)) else esc(val)}</div></div>"
        for lbl, val, css in summary_cards
    ) + "</div>"

    stmt_rows: List[str] = []
    for line in statement_lines:
        line_type = line.get("type")
        label = line.get("label_fa" if is_fa else "label_en", "")
        amount = line.get("amount", 0)
        if line_type == "section_header":
            stmt_rows.append(f"<tr class='section-row'><td colspan='4'>{esc(label)}</td></tr>")
        elif line_type == "account":
            indent = "&nbsp;" * (int(line.get("level", 0)) * 4)
            if has_compare:
                stmt_rows.append(
                    "<tr>"
                    f"<td></td><td>{esc(line.get('account_code', ''))}</td>"
                    f"<td class='name'>{indent}{esc(line.get('account_name', ''))}</td>"
                    f"<td class='num'>{fmt(amount)}</td>"
                    f"<td class='num'>{fmt(line.get('prior_amount', 0))}</td>"
                    f"<td class='num'>{fmt(line.get('variance', 0))}</td>"
                    "</tr>"
                )
            else:
                stmt_rows.append(
                    "<tr>"
                    f"<td></td><td>{esc(line.get('account_code', ''))}</td>"
                    f"<td class='name'>{indent}{esc(line.get('account_name', ''))}</td>"
                    f"<td class='num'>{fmt(amount)}</td>"
                    "</tr>"
                )
        elif line_type in ("subtotal", "grand_total"):
            css = "highlight-row" if line.get("highlight") else "subtotal-row"
            if has_compare:
                stmt_rows.append(
                    f"<tr class='{css}'>"
                    f"<td class='name' colspan='3'>{esc(label)}</td>"
                    f"<td class='num'>{fmt(amount)}</td>"
                    f"<td class='num'>{fmt(line.get('prior_amount', 0))}</td>"
                    f"<td class='num'>{fmt(line.get('variance', 0))}</td>"
                    "</tr>"
                )
            else:
                stmt_rows.append(
                    f"<tr class='{css}'>"
                    f"<td class='name' colspan='3'>{esc(label)}</td>"
                    f"<td class='num'>{fmt(amount)}</td>"
                    "</tr>"
                )
        elif line_type == "equation_check":
            css = "equation-balanced" if line.get("equation_balanced") else "equation-unbalanced"
            stmt_rows.append(
                f"<tr class='{css}'>"
                f"<td class='name' colspan='3'>{esc(label)}</td>"
                f"<td class='num'>{fmt(amount)}</td>"
                "</tr>"
            )

    col_desc = "شرح" if is_fa else "Description"
    col_code = "کد" if is_fa else "Code"
    col_name = "نام حساب" if is_fa else "Account"
    col_amount = "مبلغ" if is_fa else "Amount"
    header_cells = f"<th>{esc(col_desc)}</th><th>{esc(col_code)}</th><th>{esc(col_name)}</th><th>{esc(col_amount)}</th>"
    if has_compare:
        header_cells += f"<th>{esc('دوره قبل' if is_fa else 'Prior')}</th><th>{esc('تغییر' if is_fa else 'Change')}</th>"

    table_html = f"""
    <html lang="{html_lang}" dir="{html_dir}">
      <head><meta charset="utf-8"></head>
      <body>
        <div class="header">
          <div class="title">{esc(title_text)}</div>
          <div class="meta">
            <div><strong>{esc('نام کسب‌وکار' if is_fa else 'Business')}:</strong> {esc(business_name)}</div>
            <div><strong>{esc('تاریخ گزارش' if is_fa else 'Report Date')}:</strong> {esc(generated_at)}</div>
          </div>
        </div>
        {filters_html}
        {summary_html}
        <table class="report-table">
          <thead><tr>{header_cells}</tr></thead>
          <tbody>{''.join(stmt_rows)}</tbody>
        </table>
      </body>
    </html>
    """

    fa_font_url_regular = None
    fa_font_url_bold = None
    try:
        if is_fa:
            fa_font_url_regular, fa_font_url_bold = load_farsi_font_data_uris()
    except Exception:
        pass

    font_face_css = ""
    if fa_font_url_regular:
        font_face_css += f'@font-face {{ font-family: \'YekanBakhFaNum\'; src: url("{fa_font_url_regular}") format(\'truetype\'); font-weight: 400; }}\n'
    if fa_font_url_bold:
        font_face_css += f'@font-face {{ font-family: \'YekanBakhFaNum\'; src: url("{fa_font_url_bold}") format(\'truetype\'); font-weight: 700; }}\n'

    preferred_stack = "YekanBakhFaNum, Vazirmatn, Tahoma, Arial, sans-serif" if is_fa else "Arial, sans-serif"
    injected = (
        "<style id=\"hesabix-font-inject\">"
        + (font_face_css or "")
        + f"\nhtml, body, body * {{ font-family: {preferred_stack} !important; }}\n</style>"
    )
    if "</head>" in table_html:
        table_html = table_html.replace("</head>", injected + "</head>")
    else:
        table_html = injected + table_html

    css = CSS(string=(font_face_css or "") + f"""
      @page {{ size: A4 portrait; margin: 12mm; }}
      body {{ font-family: {preferred_stack} !important; font-size: 11px; color: #222; }}
      .header {{ display: flex; justify-content: space-between; margin-bottom: 10px;
        border-bottom: 2px solid #444; padding-bottom: 6px; }}
      .title {{ font-size: 16px; font-weight: 700; }}
      .meta {{ font-size: 11px; color: #555; }}
      .filters {{ margin: 8px 0 10px; font-size: 10.5px; display: flex; flex-wrap: wrap; gap: 6px; }}
      .summary {{ display: flex; gap: 8px; margin: 12px 0 14px; }}
      .summary-card {{ flex: 1; border: 1px solid #e2e8f0; border-radius: 8px; padding: 8px; background: #fafbfc; }}
      .summary-card .label {{ font-size: 9px; color: #64748b; }}
      .summary-card .value {{ font-size: 12px; font-weight: 700; }}
      table.report-table {{ width: 100%; border-collapse: collapse; }}
      thead th {{ background: #f0f3f7; border: 1px solid #c7cdd6; padding: 6px; text-align: center; }}
      tbody td {{ border: 1px solid #d7dde6; padding: 5px; }}
      tbody td.name {{ text-align: {'right' if is_fa else 'left'}; }}
      tbody td.num {{ text-align: center; }}
      tr.section-row td {{ background: #dce6f1; font-weight: 700; }}
      tr.subtotal-row td {{ background: #fff8e1; font-weight: 700; }}
      tr.highlight-row td {{ background: #e8f4fd; font-weight: 700; }}
      tr.equation-balanced td {{ background: #ecfdf5; color: #065f46; font-weight: 700; }}
      tr.equation-unbalanced td {{ background: #fff7ed; color: #9a3412; font-weight: 700; }}
    """)
    font_config = FontConfiguration()
    return HTML(string=table_html).write_pdf(stylesheets=[css], font_config=font_config)


def balance_sheet_excel_response(
    result: Dict[str, Any],
    *,
    db: Session,
    business_id: int,
    filename_prefix: str,
    is_fa: bool,
    export_filename_timestamp: Callable[[int], str],
) -> Response:
    biz_name = ""
    try:
        b = db.query(Business).filter(Business.id == business_id).first()
        if b is not None:
            biz_name = b.name or ""
    except Exception:
        pass
    content = build_balance_sheet_excel_bytes(result, is_fa)
    filename = f"{filename_prefix}_{_slugify(biz_name) or business_id}_{export_filename_timestamp(business_id)}.xlsx"
    return Response(
        content=content,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename={filename}",
            "Content-Length": str(len(content)),
            "Access-Control-Expose-Headers": "Content-Disposition",
        },
    )


def balance_sheet_pdf_response(
    result: Dict[str, Any],
    *,
    db: Session,
    business_id: int,
    filename_prefix: str,
    is_fa: bool,
    report_title_fa: str,
    report_title_en: str,
    generated_at: str,
    fiscal_year_id: Optional[int],
    date_from: Optional[str],
    date_to: Optional[str],
    currency_id: Optional[int],
    project_id: Optional[int],
    account_level: Optional[int],
    export_filename_timestamp: Callable[[int], str],
    load_farsi_font_data_uris: Callable[[], Tuple[Optional[str], Optional[str]]],
) -> Response:
    biz_name = ""
    try:
        b = db.query(Business).filter(Business.id == business_id).first()
        if b is not None:
            biz_name = b.name or ""
    except Exception:
        pass
    filters_html = _build_filters_html(
        db, business_id, is_fa, fiscal_year_id, date_from, date_to, currency_id, project_id, account_level
    )
    pdf_bytes = build_balance_sheet_pdf_bytes(
        result,
        is_fa=is_fa,
        business_name=biz_name,
        report_title_fa=report_title_fa,
        report_title_en=report_title_en,
        generated_at=generated_at,
        filters_html=filters_html,
        load_farsi_font_data_uris=load_farsi_font_data_uris,
    )
    filename = f"{filename_prefix}_{_slugify(biz_name) or business_id}_{export_filename_timestamp(business_id)}.pdf"
    return Response(
        content=pdf_bytes,
        media_type="application/pdf",
        headers={
            "Content-Disposition": f"attachment; filename={filename}",
            "Content-Length": str(len(pdf_bytes)),
            "Access-Control-Expose-Headers": "Content-Disposition",
        },
    )
