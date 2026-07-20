"""خروجی PDF/Excel بسته یکپارچه گزارش‌های مالی."""
from __future__ import annotations

import io
import re
from html import escape
from typing import Any, Callable, Dict, List, Optional, Tuple

from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy.orm import Session
from weasyprint import CSS, HTML
from weasyprint.text.fonts import FontConfiguration

from adapters.db.models.business import Business
from adapters.db.models.currency import Currency
from adapters.db.models.fiscal_year import FiscalYear


def _slugify(text: str) -> str:
    return re.sub(r"[^A-Za-z0-9_-]+", "_", str(text)).strip("_")


def _format_number_for_export(raw_value: Any, is_fa: bool) -> str:
    if raw_value is None:
        return "0"
    try:
        f = float(raw_value)
        if f == int(f):
            s = f"{int(f):,}"
        else:
            s = f"{f:,.2f}"
    except Exception:
        s = str(raw_value)
    if is_fa:
        s = s.replace(",", "٬")
    return s


def _build_filters_html(
    db: Session,
    *,
    is_fa: bool,
    fiscal_year_id: Optional[int],
    date_from: Optional[str],
    date_to: Optional[str],
    currency_id: Optional[int],
    project_id: Optional[int],
    account_level: Optional[int] = None,
    column_mode: Optional[int] = None,
) -> str:
    filters: List[Tuple[str, str]] = []
    if fiscal_year_id:
        fy_title = str(fiscal_year_id)
        try:
            fy_obj = db.query(FiscalYear).filter(FiscalYear.id == int(fiscal_year_id)).first()
            if fy_obj:
                fy_title = fy_obj.title or fy_title
        except Exception:
            pass
        filters.append(("سال مالی" if is_fa else "Fiscal Year", fy_title))
    if date_from:
        filters.append(("از تاریخ" if is_fa else "From", date_from))
    if date_to:
        filters.append(("تا تاریخ" if is_fa else "To", date_to))
    if currency_id:
        cur_label = str(currency_id)
        try:
            cur = db.query(Currency).filter(Currency.id == int(currency_id)).first()
            if cur is not None:
                cur_label = (cur.title or cur.name or "") if is_fa else (cur.code or cur.name or cur.title or "")
        except Exception:
            pass
        filters.append(("ارز" if is_fa else "Currency", cur_label))
    if project_id:
        filters.append(("پروژه" if is_fa else "Project", str(project_id)))
    if account_level:
        level_labels = {1: "گروه", 2: "کل", 3: "معین", 4: "تفصیل"}
        filters.append(("سطح نمایش" if is_fa else "Level", level_labels.get(int(account_level), str(account_level))))
    if column_mode:
        filters.append(("ستون‌های تراز آزمایشی" if is_fa else "TB Columns", str(column_mode)))

    if not filters:
        return ""
    parts = [
        f"<span class='filter-item'><strong>{escape(k)}:</strong> {escape(v)}</span>"
        for k, v in filters
        if v
    ]
    return f"<div class='filters'>{''.join(parts)}</div>" if parts else ""


def _statement_rows_html(statement_lines: List[Dict[str, Any]], *, is_fa: bool, has_compare: bool) -> str:
    def esc(v: Any) -> str:
        return escape("" if v is None else str(v))

    def fmt(raw_value: Any) -> str:
        return esc(_format_number_for_export(raw_value, is_fa))

    rows: List[str] = []
    for line in statement_lines:
        line_type = line.get("type")
        label = line.get("label_fa" if is_fa else "label_en", "")
        amount = line.get("amount", 0)
        if line_type == "section_header":
            colspan = 6 if has_compare else 4
            rows.append(f"<tr class='section-row'><td colspan='{colspan}'>{esc(label)}</td></tr>")
        elif line_type == "account":
            if has_compare:
                rows.append(
                    "<tr>"
                    f"<td></td><td>{esc(line.get('account_code', ''))}</td>"
                    f"<td class='name'>{esc(line.get('account_name', ''))}</td>"
                    f"<td class='num'>{fmt(amount)}</td>"
                    f"<td class='num'>{fmt(line.get('prior_amount', 0))}</td>"
                    f"<td class='num'>{fmt(line.get('variance', 0))}</td>"
                    "</tr>"
                )
            else:
                rows.append(
                    "<tr>"
                    f"<td></td><td>{esc(line.get('account_code', ''))}</td>"
                    f"<td class='name'>{esc(line.get('account_name', ''))}</td>"
                    f"<td class='num'>{fmt(amount)}</td>"
                    "</tr>"
                )
        elif line_type == "subtotal":
            css = "highlight-row" if line.get("highlight") else "subtotal-row"
            if has_compare:
                rows.append(
                    f"<tr class='{css}'>"
                    f"<td class='name' colspan='3'>{esc(label)}</td>"
                    f"<td class='num'>{fmt(amount)}</td>"
                    f"<td class='num'>{fmt(line.get('prior_amount', 0))}</td>"
                    f"<td class='num'>{fmt(line.get('variance', 0))}</td>"
                    "</tr>"
                )
            else:
                rows.append(
                    f"<tr class='{css}'>"
                    f"<td class='name' colspan='3'>{esc(label)}</td>"
                    f"<td class='num'>{fmt(amount)}</td>"
                    "</tr>"
                )
        elif line_type == "grand_total":
            net = float(amount or 0)
            net_class = "profit" if net >= 0 else "loss"
            if has_compare:
                rows.append(
                    f"<tr class='grand-total {net_class}'>"
                    f"<td class='name' colspan='3'>{esc(label)}</td>"
                    f"<td class='num'>{fmt(amount)}</td>"
                    f"<td class='num'>{fmt(line.get('prior_amount', 0))}</td>"
                    f"<td class='num'>{fmt(line.get('variance', 0))}</td>"
                    "</tr>"
                )
            else:
                rows.append(
                    f"<tr class='grand-total {net_class}'>"
                    f"<td class='name' colspan='3'>{esc(label)}</td>"
                    f"<td class='num'>{fmt(amount)}</td>"
                    "</tr>"
                )
    return "".join(rows)


def _trial_balance_rows_html(items: List[Dict[str, Any]], *, is_fa: bool, column_mode: int) -> str:
    def esc(v: Any) -> str:
        return escape("" if v is None else str(v))

    def fmt(raw_value: Any) -> str:
        return esc(_format_number_for_export(raw_value, is_fa))

    rows: List[str] = []
    for item in items:
        cells = [
            f"<td>{esc(item.get('account_code', ''))}</td>",
            f"<td class='name'>{esc(item.get('account_name', ''))}</td>",
        ]
        if column_mode >= 6:
            cells.extend([
                f"<td class='num'>{fmt(item.get('opening_debit', 0))}</td>",
                f"<td class='num'>{fmt(item.get('opening_credit', 0))}</td>",
            ])
        if column_mode >= 4:
            cells.extend([
                f"<td class='num'>{fmt(item.get('period_debit', 0))}</td>",
                f"<td class='num'>{fmt(item.get('period_credit', 0))}</td>",
            ])
        cells.extend([
            f"<td class='num'>{fmt(item.get('closing_debit', 0))}</td>",
            f"<td class='num'>{fmt(item.get('closing_credit', 0))}</td>",
        ])
        rows.append("<tr>" + "".join(cells) + "</tr>")
    return "".join(rows)


def _trial_balance_header_html(*, is_fa: bool, column_mode: int) -> str:
    headers = [
        "کد حساب" if is_fa else "Code",
        "نام حساب" if is_fa else "Account",
    ]
    if column_mode >= 6:
        headers.extend([
            "بدهکار اول دوره" if is_fa else "Open Dr",
            "بستانکار اول دوره" if is_fa else "Open Cr",
        ])
    if column_mode >= 4:
        headers.extend([
            "بدهکار دوره" if is_fa else "Period Dr",
            "بستانکار دوره" if is_fa else "Period Cr",
        ])
    headers.extend([
        "بدهکار پایان دوره" if is_fa else "Close Dr",
        "بستانکار پایان دوره" if is_fa else "Close Cr",
    ])
    return "".join(f"<th>{escape(h)}</th>" for h in headers)


def _statement_header_html(*, is_fa: bool, has_compare: bool) -> str:
    col_desc = "شرح" if is_fa else "Description"
    col_code = "کد" if is_fa else "Code"
    col_name = "نام حساب" if is_fa else "Account"
    col_amount = "مبلغ" if is_fa else "Amount"
    cells = f"<th>{escape(col_desc)}</th><th>{escape(col_code)}</th><th>{escape(col_name)}</th><th>{escape(col_amount)}</th>"
    if has_compare:
        col_prior = "دوره قبل" if is_fa else "Prior"
        col_var = "تغییر" if is_fa else "Change"
        cells += f"<th>{escape(col_prior)}</th><th>{escape(col_var)}</th>"
    return cells


def build_financial_package_pdf_bytes(
    package: Dict[str, Any],
    *,
    is_fa: bool,
    business_name: str,
    generated_at: str,
    filters_html: str,
    column_mode: int,
    load_farsi_font_data_uris: Callable[[], Tuple[Optional[str], Optional[str]]],
) -> bytes:
    trial_balance = package.get("trial_balance") or {}
    balance_sheet = package.get("balance_sheet") or {}
    pnl_period = package.get("pnl_period") or {}
    pkg_summary = package.get("package_summary") or {}

    tb_summary = trial_balance.get("summary") or {}
    bs_summary = balance_sheet.get("summary") or {}
    pnl_summary = pnl_period.get("summary") or {}
    tb_items = trial_balance.get("items") or []
    bs_lines = balance_sheet.get("statement_lines") or []
    pnl_lines = pnl_period.get("statement_lines") or []
    bs_compare = balance_sheet.get("comparison") or {}
    pnl_compare = pnl_period.get("comparison") or {}
    has_bs_compare = bool(bs_compare)
    has_pnl_compare = bool(pnl_compare)

    def esc(v: Any) -> str:
        return escape("" if v is None else str(v))

    def fmt(raw_value: Any) -> str:
        return esc(_format_number_for_export(raw_value, is_fa))

    title_text = "بسته گزارش‌های مالی" if is_fa else "Financial Reports Package"
    subtitle = "تراز آزمایشی · سود و زیان · ترازنامه" if is_fa else "Trial Balance · P&L · Balance Sheet"
    html_lang = "fa" if is_fa else "en"
    html_dir = "rtl" if is_fa else "ltr"
    page_label_left = "صفحه " if is_fa else "Page "
    page_label_of = " از " if is_fa else " of "

    tb_valid = bool(pkg_summary.get("trial_balance_valid", tb_summary.get("balance_valid")))
    bs_balanced = bool(pkg_summary.get("balance_sheet_equation_balanced", bs_summary.get("equation_balanced")))
    net_profit = float(pkg_summary.get("net_profit_after_tax", pnl_summary.get("net_profit_after_tax", 0)) or 0)
    net_class = "profit" if net_profit >= 0 else "loss"

    overview_cards = [
        (
            "وضعیت تراز آزمایشی" if is_fa else "Trial Balance",
            "متوازن" if tb_valid else "نامتوازن",
            "balanced" if tb_valid else "unbalanced",
        ),
        ("جمع دارایی‌ها" if is_fa else "Total Assets", pkg_summary.get("total_assets", bs_summary.get("total_assets", 0)), "assets"),
        ("جمع بدهی‌ها" if is_fa else "Total Liabilities", pkg_summary.get("total_liabilities", bs_summary.get("total_liabilities", 0)), "liab"),
        ("سود خالص دوره" if is_fa else "Net Profit", net_profit, net_class),
        (
            "معادله ترازنامه" if is_fa else "BS Equation",
            "متوازن" if bs_balanced else f"اختلاف: {pkg_summary.get('balance_sheet_difference', 0)}",
            "balanced" if bs_balanced else "unbalanced",
        ),
    ]
    overview_html = "<div class='summary'>" + "".join(
        f"<div class='summary-card {css}'><div class='label'>{esc(lbl)}</div>"
        f"<div class='value'>{fmt(val) if isinstance(val, (int, float)) else esc(val)}</div></div>"
        for lbl, val, css in overview_cards
    ) + "</div>"

    tb_section_title = "۱. تراز آزمایشی" if is_fa else "1. Trial Balance"
    pnl_section_title = "۲. صورت سود و زیان" if is_fa else "2. Income Statement"
    bs_section_title = "۳. ترازنامه (صورت وضعیت مالی)" if is_fa else "3. Balance Sheet"

    tb_totals_row = ""
    if tb_summary:
        tb_totals_row = (
            "<tr class='subtotal-row'>"
            f"<td colspan='2' class='name'>{'جمع کل' if is_fa else 'Totals'}</td>"
        )
        if column_mode >= 6:
            tb_totals_row += (
                f"<td class='num'>{fmt(tb_summary.get('total_opening_debit', 0))}</td>"
                f"<td class='num'>{fmt(tb_summary.get('total_opening_credit', 0))}</td>"
            )
        if column_mode >= 4:
            tb_totals_row += (
                f"<td class='num'>{fmt(tb_summary.get('total_period_debit', 0))}</td>"
                f"<td class='num'>{fmt(tb_summary.get('total_period_credit', 0))}</td>"
            )
        tb_totals_row += (
            f"<td class='num'>{fmt(tb_summary.get('total_closing_debit', 0))}</td>"
            f"<td class='num'>{fmt(tb_summary.get('total_closing_credit', 0))}</td>"
            "</tr>"
        )

    document_html = f"""
    <html lang="{html_lang}" dir="{html_dir}">
      <head><meta charset="utf-8"></head>
      <body>
        <div class="header">
          <div>
            <div class="title">{esc(title_text)}</div>
            <div class="subtitle">{esc(subtitle)}</div>
          </div>
          <div class="meta">
            <div><strong>{'نام کسب‌وکار' if is_fa else 'Business'}:</strong> {esc(business_name)}</div>
            <div><strong>{'تاریخ گزارش' if is_fa else 'Report Date'}:</strong> {esc(generated_at)}</div>
          </div>
        </div>
        {filters_html}
        {overview_html}

        <div class="report-section">
          <h2 class="section-title">{esc(tb_section_title)}</h2>
          <table class="report-table">
            <thead><tr>{_trial_balance_header_html(is_fa=is_fa, column_mode=column_mode)}</tr></thead>
            <tbody>
              {_trial_balance_rows_html(tb_items, is_fa=is_fa, column_mode=column_mode)}
              {tb_totals_row}
            </tbody>
          </table>
        </div>

        <div class="report-section page-break">
          <h2 class="section-title">{esc(pnl_section_title)}</h2>
          <table class="report-table">
            <thead><tr>{_statement_header_html(is_fa=is_fa, has_compare=has_pnl_compare)}</tr></thead>
            <tbody>{_statement_rows_html(pnl_lines, is_fa=is_fa, has_compare=has_pnl_compare)}</tbody>
          </table>
        </div>

        <div class="report-section page-break">
          <h2 class="section-title">{esc(bs_section_title)}</h2>
          <table class="report-table">
            <thead><tr>{_statement_header_html(is_fa=is_fa, has_compare=has_bs_compare)}</tr></thead>
            <tbody>{_statement_rows_html(bs_lines, is_fa=is_fa, has_compare=has_bs_compare)}</tbody>
          </table>
        </div>
      </body>
    </html>
    """

    fa_font_url_regular = None
    fa_font_url_bold = None
    try:
        if is_fa:
            fa_font_url_regular, fa_font_url_bold = load_farsi_font_data_uris()
    except Exception:
        pass

    font_face_css = ""
    if fa_font_url_regular:
        font_face_css += (
            f'@font-face {{ font-family: \'YekanBakhFaNum\'; '
            f'src: url("{fa_font_url_regular}") format(\'truetype\'); font-weight: 400; }}\n'
        )
    if fa_font_url_bold:
        font_face_css += (
            f'@font-face {{ font-family: \'YekanBakhFaNum\'; '
            f'src: url("{fa_font_url_bold}") format(\'truetype\'); font-weight: 700; }}\n'
        )

    preferred_stack = (
        "YekanBakhFaNum, Vazirmatn, Tahoma, Arial, sans-serif"
        if is_fa
        else "Arial, sans-serif"
    )
    injected = (
        "<style id=\"hesabix-font-inject\">"
        + (font_face_css or "")
        + f"\nhtml, body, body * {{ font-family: {preferred_stack} !important; }}\n</style>"
    )
    if "</head>" in document_html:
        document_html = document_html.replace("</head>", injected + "</head>")
    else:
        document_html = injected + document_html

    css = CSS(string=(font_face_css or "") + f"""
      @page {{ size: A4 portrait; margin: 12mm;
        @bottom-{'left' if is_fa else 'right'} {{
          content: "{page_label_left}" counter(page) "{page_label_of}" counter(pages);
          font-size: 10px; color: #666;
          font-family: {preferred_stack} !important;
        }}
      }}
      body {{ font-family: {preferred_stack} !important; font-size: 11px; color: #222; }}
      .header {{ display: flex; justify-content: space-between; margin-bottom: 10px;
        border-bottom: 2px solid #444; padding-bottom: 6px; gap: 12px; }}
      .title {{ font-size: 16px; font-weight: 700; }}
      .subtitle {{ font-size: 11px; color: #64748b; margin-top: 2px; }}
      .meta {{ font-size: 11px; color: #555; text-align: {'right' if is_fa else 'left'}; }}
      .filters {{ margin: 8px 0 10px; font-size: 10.5px; display: flex; flex-wrap: wrap; gap: 6px; }}
      .filters .filter-item {{ background: #f7f9fc; border: 1px solid #e2e8f0; padding: 3px 6px; border-radius: 6px; }}
      .summary {{ display: flex; flex-wrap: wrap; gap: 8px; margin: 12px 0 14px; }}
      .summary-card {{ flex: 1 1 140px; border: 1px solid #e2e8f0; border-radius: 8px; padding: 8px 10px; background: #fafbfc; }}
      .summary-card .label {{ font-size: 9px; color: #64748b; }}
      .summary-card .value {{ font-size: 12px; font-weight: 700; }}
      .summary-card.balanced .value {{ color: #15803d; }}
      .summary-card.unbalanced .value {{ color: #b91c1c; }}
      .summary-card.assets .value {{ color: #0369a1; }}
      .summary-card.liab .value {{ color: #7c3aed; }}
      .summary-card.profit .value {{ color: #1d4ed8; }}
      .summary-card.loss .value {{ color: #c2410c; }}
      .report-section {{ margin-top: 18px; }}
      .report-section.page-break {{ page-break-before: always; }}
      h2.section-title {{ font-size: 13px; margin: 0 0 8px; color: #1e3a5f; border-bottom: 1px solid #cbd5e1; padding-bottom: 4px; }}
      table.report-table {{ width: 100%; border-collapse: collapse; }}
      thead th {{ background: #f0f3f7; border: 1px solid #c7cdd6; padding: 6px; text-align: center; font-weight: 700; font-size: 10px; }}
      tbody td {{ border: 1px solid #d7dde6; padding: 5px; vertical-align: top; font-size: 10px; }}
      tbody td.name {{ text-align: {'right' if is_fa else 'left'}; }}
      tbody td.num {{ text-align: center; font-variant-numeric: tabular-nums; }}
      tr.section-row td {{ background: #dce6f1; font-weight: 700; }}
      tr.subtotal-row td {{ background: #fff8e1; font-weight: 700; }}
      tr.highlight-row td {{ background: #e8f4fd; font-weight: 700; }}
      tr.grand-total td {{ font-weight: 800; font-size: 11px; }}
      tr.grand-total.profit td {{ background: #ecfdf5; color: #065f46; }}
      tr.grand-total.loss td {{ background: #fff7ed; color: #9a3412; }}
    """)

    font_config = FontConfiguration()
    return HTML(string=document_html).write_pdf(stylesheets=[css], font_config=font_config)


def build_financial_package_excel_bytes(package: Dict[str, Any], *, is_fa: bool, column_mode: int) -> bytes:
    trial_balance = package.get("trial_balance") or {}
    balance_sheet = package.get("balance_sheet") or {}
    pnl_period = package.get("pnl_period") or {}
    pkg_summary = package.get("package_summary") or {}

    tb_summary = trial_balance.get("summary") or {}
    bs_summary = balance_sheet.get("summary") or {}
    pnl_summary = pnl_period.get("summary") or {}
    tb_items = trial_balance.get("items") or []
    bs_lines = balance_sheet.get("statement_lines") or []
    pnl_lines = pnl_period.get("statement_lines") or []

    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2E5C8A", end_color="2E5C8A", fill_type="solid")
    section_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
    subtotal_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    thin = Side(border_style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def style_header_row(ws, ncol: int) -> None:
        for col_idx in range(1, ncol + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = border
            cell.alignment = center

    def style_data_rows(ws) -> None:
        for r in range(2, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                cell = ws.cell(row=r, column=c)
                cell.border = border
                cell.alignment = center

    # Overview sheet
    ws_overview = wb.active
    ws_overview.title = "خلاصه" if is_fa else "Overview"
    ws_overview.append(["عنوان", "مقدار"] if is_fa else ["Metric", "Value"])
    style_header_row(ws_overview, 2)
    overview_rows = [
        ("وضعیت تراز آزمایشی" if is_fa else "Trial Balance", "متوازن" if pkg_summary.get("trial_balance_valid") else "نامتوازن"),
        ("جمع دارایی‌ها" if is_fa else "Total Assets", pkg_summary.get("total_assets", bs_summary.get("total_assets", 0))),
        ("جمع بدهی‌ها" if is_fa else "Total Liabilities", pkg_summary.get("total_liabilities", bs_summary.get("total_liabilities", 0))),
        ("حقوق صاحبان سهام" if is_fa else "Equity", pkg_summary.get("total_equity", bs_summary.get("total_equity", 0))),
        ("سود خالص دوره" if is_fa else "Net Profit", pkg_summary.get("net_profit_after_tax", pnl_summary.get("net_profit_after_tax", 0))),
        (
            "معادله ترازنامه" if is_fa else "BS Equation",
            "متوازن" if pkg_summary.get("balance_sheet_equation_balanced") else pkg_summary.get("balance_sheet_difference", 0),
        ),
    ]
    for row in overview_rows:
        ws_overview.append(list(row))
    style_data_rows(ws_overview)

    # Trial balance sheet
    ws_tb = wb.create_sheet("تراز آزمایشی" if is_fa else "Trial Balance")
    tb_headers = ["کد", "نام حساب"] if is_fa else ["Code", "Account"]
    if column_mode >= 6:
        tb_headers.extend(["بدهکار اول", "بستانکار اول"] if is_fa else ["Open Dr", "Open Cr"])
    if column_mode >= 4:
        tb_headers.extend(["بدهکار دوره", "بستانکار دوره"] if is_fa else ["Period Dr", "Period Cr"])
    tb_headers.extend(["بدهکار پایان", "بستانکار پایان"] if is_fa else ["Close Dr", "Close Cr"])
    ws_tb.append(tb_headers)
    style_header_row(ws_tb, len(tb_headers))
    for item in tb_items:
        row = [item.get("account_code", ""), item.get("account_name", "")]
        if column_mode >= 6:
            row.extend([item.get("opening_debit", 0), item.get("opening_credit", 0)])
        if column_mode >= 4:
            row.extend([item.get("period_debit", 0), item.get("period_credit", 0)])
        row.extend([item.get("closing_debit", 0), item.get("closing_credit", 0)])
        ws_tb.append(row)
    totals_row = ["", "جمع کل" if is_fa else "Totals"]
    if column_mode >= 6:
        totals_row.extend([tb_summary.get("total_opening_debit", 0), tb_summary.get("total_opening_credit", 0)])
    if column_mode >= 4:
        totals_row.extend([tb_summary.get("total_period_debit", 0), tb_summary.get("total_period_credit", 0)])
    totals_row.extend([tb_summary.get("total_closing_debit", 0), tb_summary.get("total_closing_credit", 0)])
    ws_tb.append(totals_row)
    style_data_rows(ws_tb)

    def append_statement_sheet(title: str, lines: List[Dict[str, Any]]) -> None:
        ws = wb.create_sheet(title)
        headers = ["شرح", "کد", "نام حساب", "مبلغ"] if is_fa else ["Description", "Code", "Account", "Amount"]
        ws.append(headers)
        style_header_row(ws, len(headers))
        for line in lines:
            line_type = line.get("type")
            if line_type == "section_header":
                ws.append([line.get("label_fa" if is_fa else "label_en", ""), "", "", ""])
                for c in range(1, 5):
                    ws.cell(row=ws.max_row, column=c).fill = section_fill
            elif line_type == "account":
                ws.append(["", line.get("account_code", ""), line.get("account_name", ""), line.get("amount", 0)])
            elif line_type in ("subtotal", "grand_total"):
                ws.append([line.get("label_fa" if is_fa else "label_en", ""), "", "", line.get("amount", 0)])
                fill = subtotal_fill if line_type == "subtotal" else PatternFill(start_color="E8F4FD", end_color="E8F4FD", fill_type="solid")
                for c in range(1, 5):
                    ws.cell(row=ws.max_row, column=c).fill = fill
        style_data_rows(ws)

    append_statement_sheet("سود و زیان" if is_fa else "P&L", pnl_lines)
    append_statement_sheet("ترازنامه" if is_fa else "Balance Sheet", bs_lines)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def financial_package_pdf_response(
    package: Dict[str, Any],
    *,
    db: Session,
    business_id: int,
    is_fa: bool,
    generated_at: str,
    fiscal_year_id: Optional[int],
    date_from: Optional[str],
    date_to: Optional[str],
    currency_id: Optional[int],
    project_id: Optional[int],
    account_level: int,
    column_mode: int,
    export_filename_timestamp: Callable[[int], str],
    load_farsi_font_data_uris: Callable[[], Tuple[Optional[str], Optional[str]]],
) -> Response:
    biz_name = ""
    try:
        b = db.query(Business).filter(Business.id == business_id).first()
        if b is not None:
            biz_name = b.name or ""
    except Exception:
        pass

    filters_html = _build_filters_html(
        db,
        is_fa=is_fa,
        fiscal_year_id=fiscal_year_id,
        date_from=date_from,
        date_to=date_to,
        currency_id=currency_id,
        project_id=project_id,
        account_level=account_level,
        column_mode=column_mode,
    )
    content = build_financial_package_pdf_bytes(
        package,
        is_fa=is_fa,
        business_name=biz_name,
        generated_at=generated_at,
        filters_html=filters_html,
        column_mode=column_mode,
        load_farsi_font_data_uris=load_farsi_font_data_uris,
    )
    filename = f"financial_package_{_slugify(biz_name) or business_id}_{export_filename_timestamp(business_id)}.pdf"
    return Response(
        content=content,
        media_type="application/pdf",
        headers={
            "Content-Disposition": f"attachment; filename={filename}",
            "Content-Length": str(len(content)),
            "Access-Control-Expose-Headers": "Content-Disposition",
        },
    )


def financial_package_excel_response(
    package: Dict[str, Any],
    *,
    db: Session,
    business_id: int,
    is_fa: bool,
    column_mode: int,
    export_filename_timestamp: Callable[[int], str],
) -> Response:
    biz_name = ""
    try:
        b = db.query(Business).filter(Business.id == business_id).first()
        if b is not None:
            biz_name = b.name or ""
    except Exception:
        pass

    content = build_financial_package_excel_bytes(package, is_fa=is_fa, column_mode=column_mode)
    filename = f"financial_package_{_slugify(biz_name) or business_id}_{export_filename_timestamp(business_id)}.xlsx"
    return Response(
        content=content,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename={filename}",
            "Content-Length": str(len(content)),
            "Access-Control-Expose-Headers": "Content-Disposition",
        },
    )
