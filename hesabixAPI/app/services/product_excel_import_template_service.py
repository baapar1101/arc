"""ساخت فایل تمپلیت ایمپورت کالا/خدمت با شیت راهنما و فهرست مراجع."""
from __future__ import annotations

import datetime
import io
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional, Sequence, Tuple

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session

from adapters.db.models.category import BusinessCategory
from adapters.db.models.currency import BusinessCurrency, Currency
from adapters.db.models.product_attribute import ProductAttribute
from adapters.db.models.tax_type import TaxType
from adapters.db.models.tax_unit import TaxUnit
from adapters.db.models.warehouse import Warehouse
from app.services.product_excel_import_spec import (
    HEADER_GROUP_COLORS,
    SAMPLE_PRODUCT_CODE,
    USER_TEMPLATE_COLUMNS,
    ImportColumn,
    localized_header,
    localized_help,
)


CATEGORY_PATH_SEPARATOR = " > "


@dataclass
class ProductImportLookups:
    category_rows: List[Dict[str, Any]] = field(default_factory=list)
    warehouses: List[Dict[str, Any]] = field(default_factory=list)
    tax_types: List[Dict[str, Any]] = field(default_factory=list)
    tax_units: List[Dict[str, Any]] = field(default_factory=list)
    attributes: List[Dict[str, Any]] = field(default_factory=list)
    currencies: List[Dict[str, Any]] = field(default_factory=list)


def _title(cat: BusinessCategory, locale: str) -> str:
    trans = cat.title_translations or {}
    preferred = locale if locale in trans else "fa"
    return (
        (trans.get(preferred) or "").strip()
        or (trans.get("fa") or "").strip()
        or (trans.get("en") or "").strip()
        or ""
    )


def category_full_path(
    cat: BusinessCategory,
    by_id: Dict[int, BusinessCategory],
    locale: str,
    separator: str = CATEGORY_PATH_SEPARATOR,
) -> str:
    parts: List[str] = []
    seen: set[int] = set()
    current: Optional[BusinessCategory] = cat
    while current is not None and current.id not in seen:
        seen.add(current.id)
        label = _title(current, locale)
        if label:
            parts.append(label)
        parent_id = current.parent_id
        current = by_id.get(parent_id) if parent_id else None
    parts.reverse()
    return separator.join(parts)


def load_product_import_lookups(db: Session, business_id: int, locale: str) -> ProductImportLookups:
    cats = (
        db.query(BusinessCategory)
        .filter(BusinessCategory.business_id == business_id)
        .order_by(BusinessCategory.sort_order.asc(), BusinessCategory.id.asc())
        .all()
    )
    by_id = {c.id: c for c in cats}
    category_rows = []
    for c in cats:
        trans = c.title_translations or {}
        category_rows.append(
            {
                "path": category_full_path(c, by_id, locale),
                "title_fa": (trans.get("fa") or "").strip(),
                "title_en": (trans.get("en") or "").strip(),
            }
        )

    warehouses = [
        {
            "code": (w.code or "").strip(),
            "name": (w.name or "").strip(),
            "is_default": bool(w.is_default),
        }
        for w in db.query(Warehouse)
        .filter(Warehouse.business_id == business_id)
        .order_by(Warehouse.code.asc())
        .all()
    ]

    tax_types = [
        {"code": (t.code or "").strip(), "title": (t.title or "").strip()}
        for t in db.query(TaxType).order_by(TaxType.code.asc()).all()
    ]
    tax_units = [
        {"code": (u.code or "").strip(), "name": (u.name or "").strip()}
        for u in db.query(TaxUnit).order_by(TaxUnit.code.asc()).all()
    ]
    attributes = [
        {"title": (a.title or "").strip(), "data_type": (a.data_type or "").strip()}
        for a in db.query(ProductAttribute)
        .filter(ProductAttribute.business_id == business_id)
        .order_by(ProductAttribute.title.asc())
        .all()
    ]
    currencies = [
        {
            "code": (c.code or "").strip(),
            "title": (c.title or c.name or "").strip(),
        }
        for c in (
            db.query(Currency)
            .join(BusinessCurrency, BusinessCurrency.currency_id == Currency.id)
            .filter(BusinessCurrency.business_id == business_id)
            .order_by(Currency.code.asc())
            .all()
        )
    ]
    return ProductImportLookups(
        category_rows=category_rows,
        warehouses=warehouses,
        tax_types=tax_types,
        tax_units=tax_units,
        attributes=attributes,
        currencies=currencies,
    )


def _thin_border() -> Border:
    side = Side(style="thin", color="BFBFBF")
    return Border(left=side, right=side, top=side, bottom=side)


def _header_fill(group: str) -> PatternFill:
    color = HEADER_GROUP_COLORS.get(group, "1F4E79")
    return PatternFill(start_color=color, end_color=color, fill_type="solid")


def _set_rtl(ws: Worksheet, enabled: bool) -> None:
    if not enabled:
        return
    try:
        ws.sheet_view.rightToLeft = True
    except Exception:
        pass


def _autosize(ws: Worksheet, min_width: int = 12, max_width: int = 42) -> None:
    for column in ws.columns:
        try:
            letter = column[0].column_letter
            max_len = 0
            for cell in column:
                if cell.value is None:
                    continue
                max_len = max(max_len, min(len(str(cell.value)), max_width))
            ws.column_dimensions[letter].width = max(min_width, min(max_len + 2, max_width))
        except Exception:
            continue


def _write_header_row(ws: Worksheet, headers: Sequence[str], fills: Sequence[PatternFill], helps: Sequence[str]) -> None:
    font = Font(bold=True, color="FFFFFF")
    align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    border = _thin_border()
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = font
        cell.fill = fills[col - 1]
        cell.alignment = align
        cell.border = border
        help_text = (helps[col - 1] or "").strip()
        if help_text:
            cell.comment = Comment(help_text, "Hesabix")
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


def _lookup_sheet(
    wb: Workbook,
    title: str,
    headers: Sequence[str],
    rows: Sequence[Sequence[Any]],
    *,
    rtl: bool,
    empty_fa: str,
    empty_en: str,
    locale: str,
) -> None:
    ws = wb.create_sheet(title)
    _set_rtl(ws, rtl)
    fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    _write_header_row(ws, headers, [fill] * len(headers), [""] * len(headers))
    if not rows:
        ws.cell(row=2, column=1, value=empty_fa if locale == "fa" else empty_en)
    else:
        for r_idx, row in enumerate(rows, start=2):
            for c_idx, val in enumerate(row, start=1):
                ws.cell(row=r_idx, column=c_idx, value=val)
    _autosize(ws)
    ws.sheet_properties.tabColor = "5D6D7E"


def _instructions_sheet(wb: Workbook, locale: str, rtl: bool) -> None:
    ws = wb.create_sheet("راهنما" if locale == "fa" else "Instructions")
    _set_rtl(ws, rtl)
    ws.sheet_properties.tabColor = "1F4E79"
    title_font = Font(bold=True, size=14, color="1F4E79")
    heading_font = Font(bold=True, size=12, color="1F4E79")
    body_align = Alignment(wrap_text=True, vertical="top")

    if locale == "fa":
        blocks = [
            ("راهنمای ایمپورت کالا و خدمات", title_font),
            (
                "فقط شیت «کالاها» را پر کنید. شیت‌های دیگر فهرست مقادیر موجود کسب‌وکار شما هستند؛ "
                "از آن‌ها کپی کنید و در شیت کالاها بچسبانید. شناسه داخلی لازم نیست.",
                None,
            ),
            ("دسته‌بندی", heading_font),
            (
                "در برنامه دسته با نام و مسیر دیده می‌شود، نه با شماره. "
                "ستون «مسیر دسته‌بندی» را مثل «مواد اولیه > پلاستیک» پر کنید. "
                "مسیر کامل را از شیت «دسته‌بندی‌ها» کپی کنید. "
                "اگر دسته وجود ندارد، در دیالوگ ایمپورت گزینه «ایجاد خودکار» را بزنید.",
                None,
            ),
            ("سلول خالی", heading_font),
            (
                "هنگام به‌روزرسانی کالاهای موجود، سلول خالی یعنی «این فیلد را تغییر نده». "
                "برای غیرفعال کردن یک گزینه صریحاً «خیر» بنویسید.",
                None,
            ),
            ("مقادیر بله/خیر", heading_font),
            ("بله، خیر، true، false، 1، 0", None),
            ("نوع و حالت موجودی", heading_font),
            ("نوع: کالا یا خدمت. حالت موجودی: فله‌ای یا یونیک.", None),
            ("ردیف نمونه", heading_font),
            (
                "ردیف دوم یک نمونه است. قبل از ایمپورت واقعی آن را ویرایش یا حذف کنید. "
                "اگر بدون تغییر آپلود شود، سیستم آن را نادیده می‌گیرد.",
                None,
            ),
            ("انبار، مالیات، ویژگی، ارز", heading_font),
            (
                "کد یا نام را از شیت مرجع مربوط کپی کنید. برای چند ویژگی، عنوان‌ها را با ویرگول جدا کنید.",
                None,
            ),
        ]
    else:
        blocks = [
            ("Product & service import guide", title_font),
            (
                "Fill only the Products sheet. Other sheets list values that already exist in your business; "
                "copy from them. Internal IDs are not required.",
                None,
            ),
            ("Categories", heading_font),
            (
                "The app shows category names and paths, not numeric IDs. "
                "Fill Category Path like \"Raw materials > Plastics\". "
                "Copy the full path from the Categories sheet. "
                "If a category is missing, choose Auto-create in the import dialog.",
                None,
            ),
            ("Empty cells", heading_font),
            (
                "When updating existing products, an empty cell means \"do not change this field\". "
                "Write No explicitly to turn a flag off.",
                None,
            ),
            ("Yes / No values", heading_font),
            ("Yes, No, true, false, 1, 0", None),
            ("Type and inventory mode", heading_font),
            ("Type: product or service. Inventory mode: bulk or unique.", None),
            ("Sample row", heading_font),
            (
                "Row 2 is a sample. Edit or delete it before a real import. "
                "If uploaded unchanged, it is skipped.",
                None,
            ),
            ("Warehouse, tax, attributes, currency", heading_font),
            (
                "Copy code or name from the matching reference sheet. "
                "Separate multiple attribute titles with commas.",
                None,
            ),
        ]

    ws.column_dimensions["A"].width = 92
    for idx, (text, font) in enumerate(blocks, start=1):
        cell = ws.cell(row=idx, column=1, value=text)
        cell.alignment = body_align
        if font is not None:
            cell.font = font
        ws.row_dimensions[idx].height = 22 if font is title_font else 48 if font is None else 24


def _sample_row(locale: str, lookups: ProductImportLookups) -> List[Any]:
    category_path = ""
    if lookups.category_rows:
        category_path = lookups.category_rows[0].get("path") or ""
    warehouse_code = lookups.warehouses[0]["code"] if lookups.warehouses else ""
    warehouse_name = lookups.warehouses[0]["name"] if lookups.warehouses else ""
    if locale == "fa":
        values = {
            "code": SAMPLE_PRODUCT_CODE,
            "name": "نمونه کالا",
            "item_type": "کالا",
            "is_active": "بله",
            "description": "ردیف نمونه — قبل از ایمپورت واقعی ویرایش یا حذف کنید",
            "general_barcodes": "6260000000001",
            "category_path": category_path or "مواد اولیه > پلاستیک",
            "main_unit": "عدد",
            "base_sales_price": 150000,
            "base_purchase_price": 120000,
            "track_inventory": "بله",
            "inventory_mode": "فله‌ای",
            "warehouse_code": warehouse_code,
            "warehouse_name": warehouse_name,
            "is_sales_taxable": "خیر",
            "is_purchase_taxable": "خیر",
        }
    else:
        values = {
            "code": SAMPLE_PRODUCT_CODE,
            "name": "Sample product",
            "item_type": "product",
            "is_active": "Yes",
            "description": "Sample row — edit or delete before a real import",
            "general_barcodes": "6260000000001",
            "category_path": category_path or "Raw materials > Plastics",
            "main_unit": "pcs",
            "base_sales_price": 150000,
            "base_purchase_price": 120000,
            "track_inventory": "Yes",
            "inventory_mode": "bulk",
            "warehouse_code": warehouse_code,
            "warehouse_name": warehouse_name,
            "is_sales_taxable": "No",
            "is_purchase_taxable": "No",
        }
    return [values.get(col.key, "") for col in USER_TEMPLATE_COLUMNS]


def _add_data_validations(ws: Worksheet, columns: Sequence[ImportColumn], locale: str, last_row: int = 2000) -> None:
    key_to_col = {col.key: idx for idx, col in enumerate(columns, start=1)}

    def _list(key: str, options: str) -> None:
        idx = key_to_col.get(key)
        if not idx:
            return
        letter = get_column_letter(idx)
        dv = DataValidation(type="list", formula1=options, allow_blank=True, showDropDown=False)
        dv.error = "مقدار مجاز را از فهرست انتخاب کنید" if locale == "fa" else "Pick a value from the list"
        dv.errorTitle = "مقدار نامعتبر" if locale == "fa" else "Invalid value"
        dv.add(f"{letter}2:{letter}{last_row}")
        ws.add_data_validation(dv)

    if locale == "fa":
        _list("item_type", '"کالا,خدمت"')
        yes_no = '"بله,خیر"'
        _list("inventory_mode", '"فله‌ای,یونیک"')
    else:
        _list("item_type", '"product,service"')
        yes_no = '"Yes,No"'
        _list("inventory_mode", '"bulk,unique"')

    for key in (
        "is_active",
        "track_inventory",
        "track_serial",
        "track_barcode",
        "auto_update_base_from_fx",
        "is_sales_taxable",
        "is_purchase_taxable",
        "is_public_catalog",
    ):
        _list(key, yes_no)


def build_products_import_workbook(lookups: ProductImportLookups, locale: str) -> Workbook:
    rtl = locale == "fa"
    wb = Workbook()
    ws = wb.active
    ws.title = "کالاها" if rtl else "Products"
    _set_rtl(ws, rtl)
    ws.sheet_properties.tabColor = "1F4E79"

    headers = [localized_header(col, locale) for col in USER_TEMPLATE_COLUMNS]
    fills = [_header_fill(col.group) for col in USER_TEMPLATE_COLUMNS]
    helps = [localized_help(col, locale) for col in USER_TEMPLATE_COLUMNS]
    _write_header_row(ws, headers, fills, helps)

    sample = _sample_row(locale, lookups)
    sample_fill = PatternFill(start_color="FFF8E1", end_color="FFF8E1", fill_type="solid")
    sample_font = Font(italic=True, color="6D4C41")
    for col, val in enumerate(sample, 1):
        cell = ws.cell(row=2, column=col, value=val if val != "" else None)
        cell.fill = sample_fill
        cell.font = sample_font
        cell.border = _thin_border()

    _add_data_validations(ws, USER_TEMPLATE_COLUMNS, locale)
    _autosize(ws, min_width=14, max_width=36)

    _instructions_sheet(wb, locale, rtl)

    if rtl:
        _lookup_sheet(
            wb,
            "دسته‌بندی‌ها",
            ["مسیر دسته‌بندی", "نام فارسی", "نام انگلیسی"],
            [[r["path"], r["title_fa"], r["title_en"]] for r in lookups.category_rows],
            rtl=rtl,
            empty_fa="هنوز دسته‌بندی‌ای تعریف نشده. مسیر را بنویسید یا هنگام ایمپورت ایجاد خودکار را انتخاب کنید.",
            empty_en="",
            locale=locale,
        )
        _lookup_sheet(
            wb,
            "انبارها",
            ["کد انبار", "نام انبار", "پیش‌فرض"],
            [[w["code"], w["name"], "بله" if w["is_default"] else ""] for w in lookups.warehouses],
            rtl=rtl,
            empty_fa="انباری تعریف نشده است.",
            empty_en="",
            locale=locale,
        )
        _lookup_sheet(
            wb,
            "انواع مالیات",
            ["کد نوع مالیات", "عنوان"],
            [[t["code"], t["title"]] for t in lookups.tax_types],
            rtl=rtl,
            empty_fa="نوع مالیاتی در سیستم ثبت نشده است.",
            empty_en="",
            locale=locale,
        )
        _lookup_sheet(
            wb,
            "واحدهای مالیاتی",
            ["کد واحد مالیاتی", "نام"],
            [[u["code"], u["name"]] for u in lookups.tax_units],
            rtl=rtl,
            empty_fa="واحد مالیاتی ثبت نشده است.",
            empty_en="",
            locale=locale,
        )
        _lookup_sheet(
            wb,
            "ویژگی‌ها",
            ["نام ویژگی", "نوع داده"],
            [[a["title"], a["data_type"]] for a in lookups.attributes],
            rtl=rtl,
            empty_fa="ویژگی‌ای تعریف نشده. نام را بنویسید یا هنگام ایمپورت ایجاد خودکار را انتخاب کنید.",
            empty_en="",
            locale=locale,
        )
        _lookup_sheet(
            wb,
            "ارزها",
            ["کد ارز", "عنوان"],
            [[c["code"], c["title"]] for c in lookups.currencies],
            rtl=rtl,
            empty_fa="ارز جانبی برای این کسب‌وکار تعریف نشده است.",
            empty_en="",
            locale=locale,
        )
    else:
        _lookup_sheet(
            wb,
            "Categories",
            ["Category Path", "Title (FA)", "Title (EN)"],
            [[r["path"], r["title_fa"], r["title_en"]] for r in lookups.category_rows],
            rtl=rtl,
            empty_fa="",
            empty_en="No categories yet. Type a path or choose Auto-create on import.",
            locale=locale,
        )
        _lookup_sheet(
            wb,
            "Warehouses",
            ["Warehouse Code", "Warehouse Name", "Default"],
            [[w["code"], w["name"], "Yes" if w["is_default"] else ""] for w in lookups.warehouses],
            rtl=rtl,
            empty_fa="",
            empty_en="No warehouses defined.",
            locale=locale,
        )
        _lookup_sheet(
            wb,
            "Tax Types",
            ["Tax Type Code", "Title"],
            [[t["code"], t["title"]] for t in lookups.tax_types],
            rtl=rtl,
            empty_fa="",
            empty_en="No tax types found.",
            locale=locale,
        )
        _lookup_sheet(
            wb,
            "Tax Units",
            ["Tax Unit Code", "Name"],
            [[u["code"], u["name"]] for u in lookups.tax_units],
            rtl=rtl,
            empty_fa="",
            empty_en="No tax units found.",
            locale=locale,
        )
        _lookup_sheet(
            wb,
            "Attributes",
            ["Attribute Title", "Data Type"],
            [[a["title"], a["data_type"]] for a in lookups.attributes],
            rtl=rtl,
            empty_fa="",
            empty_en="No attributes yet. Type titles or choose Auto-create on import.",
            locale=locale,
        )
        _lookup_sheet(
            wb,
            "Currencies",
            ["Currency Code", "Title"],
            [[c["code"], c["title"]] for c in lookups.currencies],
            rtl=rtl,
            empty_fa="",
            empty_en="No extra currencies for this business.",
            locale=locale,
        )

    return wb


def build_products_import_template(
    db: Session,
    business_id: int,
    locale: str,
) -> Tuple[bytes, str]:
    lookups = load_product_import_lookups(db, business_id, locale)
    wb = build_products_import_workbook(lookups, locale)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"products_import_template_{stamp}.xlsx"
    return buf.getvalue(), filename
