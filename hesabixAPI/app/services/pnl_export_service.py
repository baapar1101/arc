"""خروجی Excel/PDF مشترک گزارش‌های سود و زیان."""
from __future__ import annotations

import io
import re
from html import escape
from typing import Any, Callable, Dict, List, Optional, Tuple

from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy.orm import Session
from weasyprint import CSS, HTML
from weasyprint.text.fonts import FontConfiguration

from adapters.db.models.business import Business
from adapters.db.models.currency import Currency
from adapters.db.models.fiscal_year import FiscalYear


def _slugify(text: str) -> str:
    return re.sub(r"[^A-Za-z0-9_-]+", "_", str(text)).strip("_")


def _format_number_for_export(raw_value: Any, is_fa: bool) -> str:
    if raw_value is None:
        return "0"
    try:
        f = float(raw_value)
        if f == int(f):
            s = f"{int(f):,}"
        else:
            s = f"{f:,.2f}"
    except Exception:
        s = str(raw_value)
    if is_fa:
        s = s.replace(",", "٬")
    return s


def build_pnl_excel_bytes(result: Dict[str, Any], is_fa: bool) -> bytes:
    summary = result.get("summary") or {}
    statement_lines = result.get("statement_lines") or []
    sales_items = result.get("sales_items") or []
    other_income_items = result.get("other_income_items") or []
    non_operating_income_items = result.get("non_operating_income_items") or []
    cogs_items = result.get("cogs_items") or []
    operating_items = result.get("operating_expense_items") or []
    non_operating_expense_items = result.get("non_operating_expense_items") or []
    tax_items = result.get("tax_expense_items") or []
    comparison = result.get("comparison") or {}
    has_compare = bool(comparison)

    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    section_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
    subtotal_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    highlight_fill = PatternFill(start_color="E8F4FD", end_color="E8F4FD", fill_type="solid")
    profit_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    loss_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    thin = Side(border_style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    net = float(summary.get("net_profit_after_tax", summary.get("net_profit_loss")) or 0)

    # Sheet: Summary
    ws_summary = wb.active
    ws_summary.title = "خلاصه" if is_fa else "Summary"
    if has_compare:
        ws_summary.append(
            ["عنوان", "دوره جاری", "دوره قبل", "تغییر"]
            if is_fa
            else ["Label", "Current", "Prior", "Variance"]
        )
    else:
        ws_summary.append(["عنوان", "مبلغ"] if is_fa else ["Label", "Amount"])
    header_cols = 4 if has_compare else 2
    for col_idx in range(1, header_cols + 1):
        cell = ws_summary.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = center

    summary_keys = [
        ("total_sales", "جمع فروش", "Total Sales"),
        ("total_revenue", "جمع درآمد عملیاتی", "Total Operating Revenue"),
        ("gross_profit", "سود ناخالص", "Gross Profit"),
        ("operating_profit", "سود عملیاتی", "Operating Profit"),
        ("total_non_operating_income", "درآمد غیرعملیاتی", "Non-Operating Income"),
        ("total_non_operating_expense", "هزینه غیرعملیاتی", "Non-Operating Expense"),
        ("profit_before_tax", "سود قبل از مالیات", "Profit Before Tax"),
        ("total_tax_expense", "مالیات بر درآمد", "Income Tax"),
        ("net_profit_after_tax", "سود خالص پس از مالیات", "Net After Tax"),
    ]
    for key, fa, en in summary_keys:
        label = fa if is_fa else en
        if has_compare and key in comparison:
            comp = comparison[key]
            ws_summary.append([
                label,
                comp.get("current", summary.get(key, 0)),
                comp.get("prior", 0),
                comp.get("variance", 0),
            ])
        else:
            ws_summary.append([label, summary.get(key, 0)])
    for r in range(2, ws_summary.max_row + 1):
        for c in range(1, header_cols + 1):
            cell = ws_summary.cell(row=r, column=c)
            cell.border = border
            cell.alignment = center
            if r == ws_summary.max_row:
                cell.fill = profit_fill if net >= 0 else loss_fill
                cell.font = Font(bold=True)
    ws_summary.column_dimensions["A"].width = 32
    ws_summary.column_dimensions["B"].width = 18
    if has_compare:
        ws_summary.column_dimensions["C"].width = 18
        ws_summary.column_dimensions["D"].width = 18
    else:
        ws_summary.column_dimensions["B"].width = 22

    def _style_header_row(ws, headers: List[str]) -> None:
        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = border
            cell.alignment = center

    def _autosize(ws) -> None:
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if cell.value is not None:
                        max_length = max(max_length, len(str(cell.value)))
                except Exception:
                    pass
            ws.column_dimensions[column_letter].width = min(max_length + 2, 60)

    def _append_detail_sheet(
        title: str,
        items: List[Dict[str, Any]],
        is_revenue: bool,
        total_key: str,
    ) -> None:
        ws = wb.create_sheet(title)
        if is_revenue:
            headers = (
                ["کد حساب", "نام حساب", "گردش بستانکار", "گردش بدهکار", "درآمد خالص"]
                if is_fa
                else ["Account Code", "Account Name", "Credit", "Debit", "Net Revenue"]
            )
        else:
            headers = (
                ["کد حساب", "نام حساب", "گردش بدهکار", "گردش بستانکار", "هزینه خالص"]
                if is_fa
                else ["Account Code", "Account Name", "Debit", "Credit", "Net Expense"]
            )
        _style_header_row(ws, headers)
        for item in items:
            if is_revenue:
                ws.append([
                    item.get("account_code", ""),
                    item.get("account_name", ""),
                    item.get("credit", 0),
                    item.get("debit", 0),
                    item.get("revenue", item.get("amount", 0)),
                ])
            else:
                ws.append([
                    item.get("account_code", ""),
                    item.get("account_name", ""),
                    item.get("debit", 0),
                    item.get("credit", 0),
                    item.get("expense", item.get("amount", 0)),
                ])
        total_label = "جمع" if is_fa else "Total"
        ws.append(["", total_label, "", "", summary.get(total_key, 0)])
        total_row = ws.max_row
        for r in range(2, total_row + 1):
            for c in range(1, 6):
                cell = ws.cell(row=r, column=c)
                cell.border = border
                cell.alignment = center
                if r == total_row:
                    cell.fill = subtotal_fill
                    cell.font = Font(bold=True)
        _autosize(ws)

    _append_detail_sheet("فروش" if is_fa else "Sales", sales_items, True, "total_sales")
    if other_income_items:
        _append_detail_sheet(
            "درآمد عملیاتی (۶۰۱)" if is_fa else "Operating Other Income",
            other_income_items,
            True,
            "total_other_income",
        )
    if non_operating_income_items:
        _append_detail_sheet(
            "درآمد غیرعملیاتی" if is_fa else "Non-Operating Income",
            non_operating_income_items,
            True,
            "total_non_operating_income",
        )
    _append_detail_sheet("بهای تمام‌شده" if is_fa else "COGS", cogs_items, False, "total_cogs")
    _append_detail_sheet(
        "هزینه عملیاتی" if is_fa else "Operating Expenses",
        operating_items,
        False,
        "total_operating_expense",
    )
    if non_operating_expense_items:
        _append_detail_sheet(
            "هزینه غیرعملیاتی" if is_fa else "Non-Operating Expenses",
            non_operating_expense_items,
            False,
            "total_non_operating_expense",
        )
    if tax_items:
        _append_detail_sheet(
            "مالیات" if is_fa else "Income Tax",
            tax_items,
            False,
            "total_tax_expense",
        )

    # Sheet: Statement
    ws_stmt = wb.create_sheet("صورت سود و زیان" if is_fa else "PnL Statement")
    if has_compare:
        stmt_headers = (
            ["شرح", "کد حساب", "نام حساب", "دوره جاری", "دوره قبل", "تغییر"]
            if is_fa
            else ["Description", "Code", "Account", "Current", "Prior", "Variance"]
        )
    else:
        stmt_headers = (
            ["شرح", "کد حساب", "نام حساب", "مبلغ"]
            if is_fa
            else ["Description", "Account Code", "Account Name", "Amount"]
        )
    _style_header_row(ws_stmt, stmt_headers)
    for line in statement_lines:
        line_type = line.get("type")
        if line_type == "section_header":
            label = line.get("label_fa" if is_fa else "label_en", "")
            ws_stmt.append([label] + [""] * (len(stmt_headers) - 1))
            row = ws_stmt.max_row
            for c in range(1, len(stmt_headers) + 1):
                cell = ws_stmt.cell(row=row, column=c)
                cell.fill = section_fill
                cell.font = Font(bold=True)
        elif line_type == "account":
            if has_compare:
                ws_stmt.append([
                    "",
                    line.get("account_code", ""),
                    line.get("account_name", ""),
                    line.get("amount", 0),
                    line.get("prior_amount", 0),
                    line.get("variance", 0),
                ])
            else:
                ws_stmt.append([
                    "",
                    line.get("account_code", ""),
                    line.get("account_name", ""),
                    line.get("amount", 0),
                ])
        elif line_type in ("subtotal", "grand_total"):
            label = line.get("label_fa" if is_fa else "label_en", "")
            if has_compare:
                ws_stmt.append([
                    label, "", "",
                    line.get("amount", 0),
                    line.get("prior_amount", 0),
                    line.get("variance", 0),
                ])
            else:
                ws_stmt.append([label, "", "", line.get("amount", 0)])
            row = ws_stmt.max_row
            fill = highlight_fill if line.get("highlight") else subtotal_fill
            if line_type == "grand_total":
                fill = profit_fill if net >= 0 else loss_fill
            for c in range(1, len(stmt_headers) + 1):
                cell = ws_stmt.cell(row=row, column=c)
                cell.fill = fill
                cell.font = Font(bold=True)
    for r in range(2, ws_stmt.max_row + 1):
        for c in range(1, len(stmt_headers) + 1):
            ws_stmt.cell(row=r, column=c).border = border
            ws_stmt.cell(row=r, column=c).alignment = center
    _autosize(ws_stmt)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def _build_filters_html(
    db: Session,
    business_id: int,
    is_fa: bool,
    fiscal_year_id: Optional[int],
    date_from: Optional[str],
    date_to: Optional[str],
    currency_id: Optional[int],
    project_id: Optional[int],
) -> str:
    filters: List[Tuple[str, str]] = []
    if fiscal_year_id:
        fy_title = str(fiscal_year_id)
        try:
            fy_obj = (
                db.query(FiscalYear)
                .filter(
                    FiscalYear.id == int(fiscal_year_id),
                    FiscalYear.business_id == int(business_id),
                )
                .first()
            )
            if fy_obj:
                fy_title = fy_obj.title or fy_title
        except Exception:
            pass
        filters.append(("سال مالی" if is_fa else "Fiscal Year", fy_title))
    if date_from:
        filters.append(("از تاریخ" if is_fa else "From", date_from))
    if date_to:
        filters.append(("تا تاریخ" if is_fa else "To", date_to))
    if currency_id:
        cur_label = str(currency_id)
        try:
            cur = db.query(Currency).filter(Currency.id == int(currency_id)).first()
            if cur is not None:
                cur_label = (cur.title or cur.name or "") if is_fa else (cur.code or cur.name or cur.title or "")
                if cur.symbol:
                    cur_label = f"{cur_label} ({cur.symbol})" if cur_label else cur.symbol
        except Exception:
            pass
        filters.append(("ارز" if is_fa else "Currency", cur_label))
    if project_id:
        filters.append(("پروژه" if is_fa else "Project", str(project_id)))

    if not filters:
        return ""
    parts = [
        f"<span class='filter-item'><strong>{escape(k)}:</strong> {escape(v)}</span>"
        for k, v in filters
        if v
    ]
    return f"<div class='filters'>{''.join(parts)}</div>" if parts else ""


def build_pnl_pdf_bytes(
    result: Dict[str, Any],
    *,
    is_fa: bool,
    business_name: str,
    report_title_fa: str,
    report_title_en: str,
    generated_at: str,
    filters_html: str,
    load_farsi_font_data_uris: Callable[[], Tuple[Optional[str], Optional[str]]],
) -> bytes:
    summary = result.get("summary") or {}
    statement_lines = result.get("statement_lines") or []
    comparison = result.get("comparison") or {}
    has_compare = bool(comparison)
    net = float(summary.get("net_profit_after_tax", summary.get("net_profit_loss")) or 0)

    def esc(v: Any) -> str:
        return escape("" if v is None else str(v))

    def fmt(raw_value: Any) -> str:
        return esc(_format_number_for_export(raw_value, is_fa))

    title_text = report_title_fa if is_fa else report_title_en
    label_biz = "نام کسب‌وکار" if is_fa else "Business Name"
    label_date = "تاریخ گزارش" if is_fa else "Report Date"
    html_lang = "fa" if is_fa else "en"
    html_dir = "rtl" if is_fa else "ltr"
    page_label_left = "صفحه " if is_fa else "Page "
    page_label_of = " از " if is_fa else " of "
    col_desc = "شرح" if is_fa else "Description"
    col_code = "کد" if is_fa else "Code"
    col_name = "نام حساب" if is_fa else "Account"
    col_amount = "مبلغ" if is_fa else "Amount"
    net_class = "profit" if net >= 0 else "loss"

    summary_cards = [
        ("جمع درآمد" if is_fa else "Total Revenue", summary.get("total_revenue", 0), "revenue"),
        ("سود قبل از مالیات" if is_fa else "Profit Before Tax", summary.get("profit_before_tax", 0), "gross"),
        ("مالیات" if is_fa else "Income Tax", summary.get("total_tax_expense", 0), "tax"),
        ("سود خالص" if is_fa else "Net After Tax", summary.get("net_profit_after_tax", summary.get("net_profit_loss", 0)), net_class),
    ]
    summary_html = "<div class='summary'>" + "".join(
        f"<div class='summary-card {css}'><div class='label'>{esc(lbl)}</div>"
        f"<div class='value'>{fmt(val)}</div></div>"
        for lbl, val, css in summary_cards
    ) + "</div>"

    stmt_rows: List[str] = []
    for line in statement_lines:
        line_type = line.get("type")
        label = line.get("label_fa" if is_fa else "label_en", "")
        amount = line.get("amount", 0)
        if line_type == "section_header":
            stmt_rows.append(
                f"<tr class='section-row'><td colspan='4'>{esc(label)}</td></tr>"
            )
        elif line_type == "account":
            if has_compare:
                stmt_rows.append(
                    "<tr>"
                    f"<td></td><td>{esc(line.get('account_code', ''))}</td>"
                    f"<td class='name'>{esc(line.get('account_name', ''))}</td>"
                    f"<td class='num'>{fmt(amount)}</td>"
                    f"<td class='num'>{fmt(line.get('prior_amount', 0))}</td>"
                    f"<td class='num'>{fmt(line.get('variance', 0))}</td>"
                    "</tr>"
                )
            else:
                stmt_rows.append(
                    "<tr>"
                    f"<td></td>"
                    f"<td>{esc(line.get('account_code', ''))}</td>"
                    f"<td class='name'>{esc(line.get('account_name', ''))}</td>"
                    f"<td class='num'>{fmt(amount)}</td>"
                    "</tr>"
                )
        elif line_type == "subtotal":
            css = "highlight-row" if line.get("highlight") else "subtotal-row"
            if has_compare:
                stmt_rows.append(
                    f"<tr class='{css}'>"
                    f"<td class='name' colspan='3'>{esc(label)}</td>"
                    f"<td class='num'>{fmt(amount)}</td>"
                    f"<td class='num'>{fmt(line.get('prior_amount', 0))}</td>"
                    f"<td class='num'>{fmt(line.get('variance', 0))}</td>"
                    "</tr>"
                )
            else:
                stmt_rows.append(
                    f"<tr class='{css}'>"
                    f"<td class='name' colspan='3'>{esc(label)}</td>"
                    f"<td class='num'>{fmt(amount)}</td>"
                    "</tr>"
                )
        elif line_type == "grand_total":
            if has_compare:
                stmt_rows.append(
                    f"<tr class='grand-total {net_class}'>"
                    f"<td class='name' colspan='3'>{esc(label)}</td>"
                    f"<td class='num'>{fmt(amount)}</td>"
                    f"<td class='num'>{fmt(line.get('prior_amount', 0))}</td>"
                    f"<td class='num'>{fmt(line.get('variance', 0))}</td>"
                    "</tr>"
                )
            else:
                stmt_rows.append(
                    f"<tr class='grand-total {net_class}'>"
                    f"<td class='name' colspan='3'>{esc(label)}</td>"
                    f"<td class='num'>{fmt(amount)}</td>"
                    "</tr>"
                )

    col_prior = "دوره قبل" if is_fa else "Prior"
    col_var = "تغییر" if is_fa else "Change"
    header_cells = f"<th>{esc(col_desc)}</th><th>{esc(col_code)}</th><th>{esc(col_name)}</th><th>{esc(col_amount)}</th>"
    if has_compare:
        header_cells += f"<th>{esc(col_prior)}</th><th>{esc(col_var)}</th>"

    table_html = f"""
    <html lang="{html_lang}" dir="{html_dir}">
      <head><meta charset="utf-8"></head>
      <body>
        <div class="header">
          <div class="title">{esc(title_text)}</div>
          <div class="meta">
            <div><strong>{esc(label_biz)}:</strong> {esc(business_name)}</div>
            <div><strong>{esc(label_date)}:</strong> {esc(generated_at)}</div>
          </div>
        </div>
        {filters_html}
        {summary_html}
        <table class="report-table">
          <thead>
            <tr>{header_cells}</tr>
          </thead>
          <tbody>{''.join(stmt_rows)}</tbody>
        </table>
      </body>
    </html>
    """

    fa_font_url_regular = None
    fa_font_url_bold = None
    try:
        if is_fa:
            fa_font_url_regular, fa_font_url_bold = load_farsi_font_data_uris()
    except Exception:
        pass

    font_face_css = ""
    if fa_font_url_regular:
        font_face_css += (
            f'@font-face {{ font-family: \'YekanBakhFaNum\'; '
            f'src: url("{fa_font_url_regular}") format(\'truetype\'); font-weight: 400; }}\n'
        )
    if fa_font_url_bold:
        font_face_css += (
            f'@font-face {{ font-family: \'YekanBakhFaNum\'; '
            f'src: url("{fa_font_url_bold}") format(\'truetype\'); font-weight: 700; }}\n'
        )

    preferred_stack = (
        "YekanBakhFaNum, Vazirmatn, Tahoma, Arial, sans-serif"
        if is_fa
        else "Arial, sans-serif"
    )
    injected = (
        "<style id=\"hesabix-font-inject\">"
        + (font_face_css or "")
        + f"\nhtml, body, body * {{ font-family: {preferred_stack} !important; }}\n</style>"
    )
    if "</head>" in table_html:
        table_html = table_html.replace("</head>", injected + "</head>")
    else:
        table_html = injected + table_html

    css = CSS(string=(font_face_css or "") + f"""
      @page {{ size: A4 portrait; margin: 12mm;
        @bottom-{'left' if is_fa else 'right'} {{
          content: "{page_label_left}" counter(page) "{page_label_of}" counter(pages);
          font-size: 10px; color: #666;
          font-family: {preferred_stack} !important;
        }}
      }}
      body {{ font-family: {preferred_stack} !important; font-size: 11px; color: #222; }}
      .header {{ display: flex; justify-content: space-between; margin-bottom: 10px;
        border-bottom: 2px solid #444; padding-bottom: 6px; gap: 12px; }}
      .title {{ font-size: 16px; font-weight: 700; }}
      .meta {{ font-size: 11px; color: #555; text-align: {'right' if is_fa else 'left'}; }}
      .filters {{ margin: 8px 0 10px; font-size: 10.5px; display: flex; flex-wrap: wrap; gap: 6px; }}
      .filters .filter-item {{ background: #f7f9fc; border: 1px solid #e2e8f0; padding: 3px 6px; border-radius: 6px; }}
      .summary {{ display: flex; gap: 8px; margin: 12px 0 14px; }}
      .summary-card {{ flex: 1; border: 1px solid #e2e8f0; border-radius: 8px; padding: 8px 10px; background: #fafbfc; }}
      .summary-card .label {{ font-size: 9px; color: #64748b; }}
      .summary-card .value {{ font-size: 13px; font-weight: 700; }}
      .summary-card.revenue .value {{ color: #15803d; }}
      .summary-card.gross .value {{ color: #0369a1; }}
      .summary-card.operating .value {{ color: #7c3aed; }}
      .summary-card.profit .value {{ color: #1d4ed8; }}
      .summary-card.loss .value {{ color: #c2410c; }}
      table.report-table {{ width: 100%; border-collapse: collapse; }}
      thead th {{ background: #f0f3f7; border: 1px solid #c7cdd6; padding: 6px; text-align: center; font-weight: 700; }}
      tbody td {{ border: 1px solid #d7dde6; padding: 5px; vertical-align: top; }}
      tbody td.name {{ text-align: {'right' if is_fa else 'left'}; }}
      tbody td.num {{ text-align: center; font-variant-numeric: tabular-nums; }}
      tr.section-row td {{ background: #dce6f1; font-weight: 700; }}
      tr.subtotal-row td {{ background: #fff8e1; font-weight: 700; }}
      tr.highlight-row td {{ background: #e8f4fd; font-weight: 700; }}
      tr.grand-total td {{ font-weight: 800; font-size: 12px; }}
      tr.grand-total.profit td {{ background: #ecfdf5; color: #065f46; }}
      tr.grand-total.loss td {{ background: #fff7ed; color: #9a3412; }}
    """)

    font_config = FontConfiguration()
    return HTML(string=table_html).write_pdf(stylesheets=[css], font_config=font_config)


def pnl_excel_response(
    result: Dict[str, Any],
    *,
    db: Session,
    business_id: int,
    filename_prefix: str,
    is_fa: bool,
    export_filename_timestamp: Callable[[int], str],
) -> Response:
    biz_name = ""
    try:
        b = db.query(Business).filter(Business.id == business_id).first()
        if b is not None:
            biz_name = b.name or ""
    except Exception:
        pass

    content = build_pnl_excel_bytes(result, is_fa)
    filename = f"{filename_prefix}_{_slugify(biz_name) or business_id}_{export_filename_timestamp(business_id)}.xlsx"
    return Response(
        content=content,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename={filename}",
            "Content-Length": str(len(content)),
            "Access-Control-Expose-Headers": "Content-Disposition",
        },
    )


def pnl_pdf_response(
    result: Dict[str, Any],
    *,
    db: Session,
    business_id: int,
    filename_prefix: str,
    is_fa: bool,
    report_title_fa: str,
    report_title_en: str,
    generated_at: str,
    fiscal_year_id: Optional[int],
    date_from: Optional[str],
    date_to: Optional[str],
    currency_id: Optional[int],
    project_id: Optional[int],
    export_filename_timestamp: Callable[[int], str],
    load_farsi_font_data_uris: Callable[[], Tuple[Optional[str], Optional[str]]],
) -> Response:
    biz_name = ""
    try:
        b = db.query(Business).filter(Business.id == business_id).first()
        if b is not None:
            biz_name = b.name or ""
    except Exception:
        pass

    filters_html = _build_filters_html(
        db, business_id, is_fa, fiscal_year_id, date_from, date_to, currency_id, project_id
    )
    pdf_bytes = build_pnl_pdf_bytes(
        result,
        is_fa=is_fa,
        business_name=biz_name,
        report_title_fa=report_title_fa,
        report_title_en=report_title_en,
        generated_at=generated_at,
        filters_html=filters_html,
        load_farsi_font_data_uris=load_farsi_font_data_uris,
    )
    filename = f"{filename_prefix}_{_slugify(biz_name) or business_id}_{export_filename_timestamp(business_id)}.pdf"
    return Response(
        content=pdf_bytes,
        media_type="application/pdf",
        headers={
            "Content-Disposition": f"attachment; filename={filename}",
            "Content-Length": str(len(pdf_bytes)),
            "Access-Control-Expose-Headers": "Content-Disposition",
        },
    )
