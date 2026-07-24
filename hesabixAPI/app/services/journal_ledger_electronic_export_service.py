"""خروجی دفتر روزنامه الکترونیکی مطابق قالب سازمان امور مالیاتی."""
from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from typing import Any, Dict, List, Literal, Optional, Tuple

from fastapi.responses import Response
from openpyxl import Workbook
from sqlalchemy.orm import Session

from adapters.db.models.business import Business
from adapters.db.models.currency import Currency
from app.core.calendar import CalendarConverter
from app.core.datetime_utils import export_filename_timestamp
from app.core.responses import ApiError
from app.services.journal_ledger_service import get_journal_ledger_report

ElectronicExportFormat = Literal["xlsx", "csv", "auto"]

# عناوین دقیق ستون‌ها مطابق قالب رسمی (شامل فاصله انتهایی «تاریخ »)
ELECTRONIC_JOURNAL_HEADERS: Tuple[str, ...] = (
    "ردیف",
    "تاریخ ",
    "کد حساب کل",
    "عنوان حساب کل",
    "کد حساب معین",
    "عنوان حساب معین",
    "شرح",
    "مبلغ بدهکار (ریال)",
    "مبلغ بستانکار (ریال)",
)

CSV_ROW_THRESHOLD = 1000
RIAL_CURRENCY_CODES = frozenset({"IRR", "RIAL"})


def _slugify(text: str) -> str:
    return re.sub(r"[^A-Za-z0-9_-]+", "_", str(text)).strip("_")


def _parse_document_date(raw_value: Any) -> Optional[date]:
    if raw_value is None:
        return None
    if isinstance(raw_value, date) and not isinstance(raw_value, datetime):
        return raw_value
    if isinstance(raw_value, datetime):
        return raw_value.date()
    if isinstance(raw_value, str):
        value = raw_value.strip()
        if not value:
            return None
        if "T" in value:
            return datetime.fromisoformat(value.replace("Z", "+00:00")).date()
        if "/" in value and len(value.split("/")) == 3:
            parts = value.split("/")
            try:
                year, month, day = int(parts[0]), int(parts[1]), int(parts[2])
                if year >= 1300:
                    import jdatetime

                    return jdatetime.date(year, month, day).togregorian()
            except (TypeError, ValueError):
                pass
        try:
            return date.fromisoformat(value.split(" ")[0])
        except ValueError:
            return None
    return None


def format_jalali_date_for_electronic_books(raw_value: Any) -> str:
    """تاریخ شمسی به فرمت YYYY/MM/DD برای سامانه دفاتر الکترونیکی."""
    doc_date = _parse_document_date(raw_value)
    if doc_date is None:
        return ""
    dt_value = datetime.combine(doc_date, datetime.min.time())
    formatted = CalendarConverter.to_jalali(dt_value)
    return formatted.get("date_only") or ""


def round_rial_amount(raw_value: Any) -> int:
    """گرد کردن مبلغ به ریال (عدد صحیح) با ROUND_HALF_UP."""
    if raw_value is None or raw_value == "":
        return 0
    amount = Decimal(str(raw_value))
    return int(amount.quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def _amount_cell_value(raw_value: Any) -> Optional[int]:
    """مقدار سلول مبلغ: فقط اگر مثبت باشد (طبق قاعده یک‌طرفه بودن سطر)."""
    rial = round_rial_amount(raw_value)
    return rial if rial > 0 else None


def item_to_electronic_row(item: Dict[str, Any], row_number: int) -> List[Any]:
    debit_value = _amount_cell_value(item.get("debit_amount"))
    credit_value = _amount_cell_value(item.get("credit_amount"))

    # سطر بدون گردش مالی در دفتر روزنامه نباید در خروجی قانونی باشد
    if debit_value is None and credit_value is None:
        return []

    return [
        row_number,
        format_jalali_date_for_electronic_books(item.get("document_date")),
        item.get("general_account_code") or "",
        item.get("general_account_name") or "",
        item.get("subsidiary_account_code") or "",
        item.get("subsidiary_account_name") or "",
        item.get("description") or "",
        debit_value,
        credit_value,
    ]


def build_electronic_journal_rows(items: List[Dict[str, Any]]) -> List[List[Any]]:
    rows: List[List[Any]] = []
    row_number = 1
    for item in items:
        row = item_to_electronic_row(item, row_number)
        if not row:
            continue
        rows.append(row)
        row_number += 1
    return rows


def _resolve_export_format(
    requested_format: str,
    row_count: int,
) -> Literal["xlsx", "csv"]:
    fmt = (requested_format or "auto").strip().lower()
    if fmt == "auto":
        return "csv" if row_count > CSV_ROW_THRESHOLD else "xlsx"
    if fmt not in ("xlsx", "csv"):
        raise ApiError(
            "INVALID_FORMAT",
            "فرمت خروجی باید xlsx، csv یا auto باشد",
            http_status=400,
        )
    if fmt == "xlsx" and row_count > CSV_ROW_THRESHOLD:
        # طبق دستورالعمل سامانه، فایل‌های بزرگ باید CSV باشند
        return "csv"
    return fmt  # type: ignore[return-value]


def _ensure_rial_currency(db: Session, currency_id: Optional[int]) -> None:
    if currency_id is None:
        raise ApiError(
            "CURRENCY_REQUIRED",
            "برای خروجی دفتر الکترونیکی باید ارز ریال (IRR) انتخاب شود",
            http_status=400,
        )
    currency = db.query(Currency).filter(Currency.id == int(currency_id)).first()
    if currency is None:
        raise ApiError("CURRENCY_NOT_FOUND", "ارز انتخاب‌شده یافت نشد", http_status=404)
    code = (currency.code or "").strip().upper()
    if code not in RIAL_CURRENCY_CODES:
        raise ApiError(
            "NON_RIAL_CURRENCY",
            "دفتر الکترونیکی باید بر اساس مبالغ ریالی باشد. لطفاً ارز ریال (IRR) را انتخاب کنید.",
            http_status=400,
        )


def build_electronic_journal_excel_bytes(rows: List[List[Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    try:
        ws.sheet_view.rightToLeft = True
    except Exception:
        pass

    for col_idx, header in enumerate(ELECTRONIC_JOURNAL_HEADERS, 1):
        ws.cell(row=1, column=col_idx, value=header)

    for row_idx, row in enumerate(rows, 2):
        for col_idx, value in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if col_idx in (8, 9) and isinstance(value, int):
                cell.number_format = "0"
            if col_idx == 2 and value:
                cell.number_format = "0"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def build_electronic_journal_csv_bytes(rows: List[List[Any]]) -> bytes:
    output = io.StringIO()
    writer = csv.writer(output, lineterminator="\n")
    writer.writerow(ELECTRONIC_JOURNAL_HEADERS)
    for row in rows:
        csv_row: List[Any] = []
        for idx, value in enumerate(row):
            if idx in (7, 8):
                csv_row.append("" if value is None else int(value))
            else:
                csv_row.append("" if value is None else value)
        writer.writerow(csv_row)
    return output.getvalue().encode("utf-8-sig")


def export_electronic_journal_books(
    db: Session,
    *,
    business_id: int,
    fiscal_year_id: Optional[int],
    currency_id: Optional[int],
    date_from: Optional[str],
    date_to: Optional[str],
    document_type: Optional[str],
    include_proforma: bool,
    export_format: str = "auto",
) -> Response:
    _ensure_rial_currency(db, currency_id)

    result = get_journal_ledger_report(
        db=db,
        business_id=business_id,
        fiscal_year_id=fiscal_year_id,
        currency_id=currency_id,
        date_from=date_from,
        date_to=date_to,
        document_type=document_type,
        include_proforma=include_proforma,
        skip=0,
        take=1,
        disable_pagination=True,
    )

    summary = result.get("summary") or {}
    if summary.get("balance_valid") is False:
        raise ApiError(
            "UNBALANCED_LEDGER",
            "دفتر روزنامه تراز نیست؛ قبل از ارسال دفتر الکترونیکی، مغایرت‌ها را برطرف کنید.",
            http_status=400,
        )

    items = result.get("items") or []
    rows = build_electronic_journal_rows(items)
    if not rows:
        raise ApiError(
            "NO_DATA",
            "در بازه انتخاب‌شده سطری برای خروجی دفتر الکترونیکی یافت نشد",
            http_status=404,
        )

    resolved_format = _resolve_export_format(export_format, len(rows))

    biz_name = ""
    try:
        business = db.query(Business).filter(Business.id == business_id).first()
        if business is not None:
            biz_name = business.name or ""
    except Exception:
        biz_name = ""

    base = "electronic_journal_ledger"
    if biz_name:
        base += f"_{_slugify(biz_name)}"
    timestamp = export_filename_timestamp(business_id)

    if resolved_format == "csv":
        content = build_electronic_journal_csv_bytes(rows)
        filename = f"{base}_{timestamp}.csv"
        media_type = "text/csv; charset=utf-8"
    else:
        content = build_electronic_journal_excel_bytes(rows)
        filename = f"{base}_{timestamp}.xlsx"
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    return Response(
        content=content,
        media_type=media_type,
        headers={
            "Content-Disposition": f"attachment; filename={filename}",
            "Content-Length": str(len(content)),
            "Access-Control-Expose-Headers": "Content-Disposition",
            "X-Export-Format": resolved_format,
            "X-Export-Row-Count": str(len(rows)),
        },
    )
