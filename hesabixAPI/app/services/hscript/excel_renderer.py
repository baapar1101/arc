"""خروجی Excel امن از Report Spec (بدون اجرای مجدد قالب‌های خطرناک)."""

from __future__ import annotations

import io
import re
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.services.hscript.errors import TypeErrorHS
from app.services.hscript.report_builder import ALLOWED_BLOCK_TYPES


def _slug_sheet(name: str, used: set[str]) -> str:
	base = re.sub(r"[\[\]\*\:\?\/\\]", "", str(name or "Sheet")).strip() or "Sheet"
	base = base[:28]
	candidate = base
	i = 2
	while candidate in used:
		suffix = f"_{i}"
		candidate = (base[: 31 - len(suffix)] + suffix)
		i += 1
	used.add(candidate)
	return candidate


def _cell(value: Any) -> Any:
	if value is None:
		return ""
	if isinstance(value, bool):
		return value
	if isinstance(value, (int, float)):
		return value
	return str(value)[:5000]


def spec_to_excel_bytes(spec: dict[str, Any], *, business_name: str | None = None) -> bytes:
	"""تبدیل Spec به workbook: خلاصه KPI + یک شیت برای هر جدول."""
	if not isinstance(spec, dict):
		raise TypeErrorHS("spec باید dict باشد")
	blocks = spec.get("blocks") or []
	if not isinstance(blocks, list):
		raise TypeErrorHS("blocks نامعتبر است")

	wb = Workbook()
	# شیت خلاصه
	ws_sum = wb.active
	ws_sum.title = "خلاصه"
	header_font = Font(bold=True, color="FFFFFF")
	header_fill = PatternFill("solid", fgColor="2563EB")
	kpi_fill = PatternFill("solid", fgColor="EFF6FF")
	thin = Border(
		left=Side(style="thin", color="D1D5DB"),
		right=Side(style="thin", color="D1D5DB"),
		top=Side(style="thin", color="D1D5DB"),
		bottom=Side(style="thin", color="D1D5DB"),
	)
	center = Alignment(horizontal="center", vertical="center", wrap_text=True)

	title = str(spec.get("title") or "گزارش HScript")[:200]
	ws_sum["A1"] = title
	ws_sum["A1"].font = Font(bold=True, size=14)
	ws_sum.merge_cells("A1:D1")
	row = 2
	if business_name:
		ws_sum[f"A{row}"] = str(business_name)[:200]
		row += 1
	layout = str(spec.get("layout") or "document")
	ws_sum[f"A{row}"] = f"layout: {layout}"
	row += 2

	ws_sum[f"A{row}"] = "نوع"
	ws_sum[f"B{row}"] = "برچسب"
	ws_sum[f"C{row}"] = "مقدار"
	ws_sum[f"D{row}"] = "توضیح"
	for col in ("A", "B", "C", "D"):
		cell = ws_sum[f"{col}{row}"]
		cell.font = header_font
		cell.fill = header_fill
		cell.alignment = center
		cell.border = thin
	row += 1

	kpis = 0
	tables: list[dict[str, Any]] = []
	charts: list[dict[str, Any]] = []

	for block in blocks:
		if not isinstance(block, dict):
			continue
		btype = str(block.get("type") or "")
		if btype not in ALLOWED_BLOCK_TYPES:
			continue
		if btype in ("kpi", "card"):
			label = block.get("label") or block.get("title") or ""
			ws_sum[f"A{row}"] = btype
			ws_sum[f"B{row}"] = _cell(label)
			ws_sum[f"C{row}"] = _cell(block.get("value"))
			ws_sum[f"D{row}"] = _cell(block.get("hint") or block.get("subtitle"))
			for col in ("A", "B", "C", "D"):
				ws_sum[f"{col}{row}"].fill = kpi_fill
				ws_sum[f"{col}{row}"].border = thin
			row += 1
			kpis += 1
		elif btype == "table":
			tables.append(block)
		elif btype == "chart":
			charts.append(block)
		elif btype in ("heading", "section", "title", "text", "paragraph"):
			ws_sum[f"A{row}"] = btype
			ws_sum[f"B{row}"] = _cell(block.get("text") or block.get("title"))
			row += 1

	if kpis == 0 and not tables and not charts:
		ws_sum[f"A{row}"] = "بلوک جدولی/آماری برای خروجی یافت نشد"
		row += 1

	for col_idx in range(1, 5):
		ws_sum.column_dimensions[get_column_letter(col_idx)].width = 22

	used_names = {"خلاصه"}

	# شیت‌های جدول
	for idx, table in enumerate(tables[:20], start=1):
		sheet_title = _slug_sheet(str(table.get("title") or f"جدول_{idx}"), used_names)
		ws = wb.create_sheet(title=sheet_title)
		cols = [str(c) for c in (table.get("columns") or [])][:50]
		rows_data = table.get("rows") or []
		if not isinstance(rows_data, list):
			rows_data = []
		if not cols and rows_data and isinstance(rows_data[0], dict):
			cols = [str(k) for k in rows_data[0].keys()][:50]
		# header
		for c_i, c in enumerate(cols, start=1):
			cell = ws.cell(row=1, column=c_i, value=c)
			cell.font = header_font
			cell.fill = header_fill
			cell.alignment = center
			cell.border = thin
		for r_i, row_dict in enumerate(rows_data[:5000], start=2):
			if not isinstance(row_dict, dict):
				continue
			for c_i, c in enumerate(cols, start=1):
				cell = ws.cell(row=r_i, column=c_i, value=_cell(row_dict.get(c)))
				cell.border = thin
		for c_i in range(1, max(len(cols), 1) + 1):
			ws.column_dimensions[get_column_letter(c_i)].width = 18

	# شیت داده نمودارها (برای تحلیل در اکسل)
	if charts:
		ws_c = wb.create_sheet(title=_slug_sheet("نمودارها", used_names))
		ws_c.append(["نمودار", "نوع", "برچسب/X", "مقدار/Y"])
		for cell in ws_c[1]:
			cell.font = header_font
			cell.fill = header_fill
			cell.border = thin
		for ch in charts[:30]:
			name = str(ch.get("title") or ch.get("chart_type") or "chart")[:80]
			ctype = str(ch.get("chart_type") or "")
			data = ch.get("data") or []
			if not isinstance(data, list):
				continue
			for item in data[:500]:
				if not isinstance(item, dict):
					continue
				if "label" in item:
					ws_c.append([name, ctype, _cell(item.get("label")), _cell(item.get("value"))])
				else:
					ws_c.append([name, ctype, _cell(item.get("x")), _cell(item.get("y"))])
		for c_i in range(1, 5):
			ws_c.column_dimensions[get_column_letter(c_i)].width = 20

	bio = io.BytesIO()
	wb.save(bio)
	return bio.getvalue()
