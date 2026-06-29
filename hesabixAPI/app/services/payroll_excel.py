"""
ورود و خروج Excel برای افزونه حقوق و دستمزد (فاز ۲).
"""
from __future__ import annotations

import datetime
import io
from decimal import Decimal, InvalidOperation
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import or_
from sqlalchemy.orm import Session

from adapters.db.models.payroll import (
	PayrollDepartment,
	PayrollEmployee,
	PayrollItemDefinition,
	PayrollRun,
	PayrollRunLine,
)
from adapters.db.models.person import Person
from app.core.responses import ApiError
from app.services.payroll_service import (
	_EMPLOYMENT_TYPES,
	_decimal,
	_ensure_plugin,
	_ensure_run_editable,
	_get_run_row,
	_persist_run_lines,
	_get_or_create_settings,
	create_employee,
	update_employee,
)

_HEADER_FILL = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _validate_xlsx(content: bytes) -> None:
	if len(content) < 100 or content[:2] != b"PK":
		raise ApiError("INVALID_FILE", "فایل Excel معتبر نیست.", http_status=400)


def _cell_str(value: Any) -> str:
	if value is None:
		return ""
	return str(value).strip()


def _cell_decimal(value: Any) -> Optional[Decimal]:
	s = _cell_str(value)
	if not s:
		return None
	try:
		return Decimal(s.replace(",", ""))
	except (InvalidOperation, ValueError):
		raise ApiError("INVALID_AMOUNT", f"مقدار عددی نامعتبر: {value}", http_status=400)


def _write_header_row(ws, headers: List[str]) -> None:
	for col, title in enumerate(headers, start=1):
		cell = ws.cell(row=1, column=col, value=title)
		cell.fill = _HEADER_FILL
		cell.font = _HEADER_FONT
		cell.alignment = _HEADER_ALIGN


def build_employees_import_template(db: Session, business_id: int) -> Tuple[bytes, str]:
	_ensure_plugin(db, business_id)
	wb = Workbook()
	ws = wb.active
	ws.title = "employees"
	headers = [
		"person_id",
		"person_code",
		"employee_code",
		"job_title",
		"base_salary",
		"department_code",
		"employment_type",
		"insurance_number",
		"tax_id",
		"hire_date",
	]
	_write_header_row(ws, headers)
	ws.append([101, "P-001", "EMP001", "حسابدار", 50000000, "finance", "full_time", "", "", "1403-01-01"])
	for col in range(1, len(headers) + 1):
		ws.column_dimensions[chr(64 + col)].width = 16
	buf = io.BytesIO()
	wb.save(buf)
	filename = f"payroll_employees_template_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
	return buf.getvalue(), filename


def export_employees_excel(db: Session, business_id: int) -> Tuple[bytes, str]:
	_ensure_plugin(db, business_id)
	rows = (
		db.query(PayrollEmployee, Person, PayrollDepartment)
		.join(Person, Person.id == PayrollEmployee.person_id)
		.outerjoin(PayrollDepartment, PayrollDepartment.id == PayrollEmployee.department_id)
		.filter(PayrollEmployee.business_id == business_id)
		.order_by(PayrollEmployee.employee_code)
		.all()
	)
	wb = Workbook()
	ws = wb.active
	ws.title = "employees"
	headers = [
		"person_id",
		"person_code",
		"employee_code",
		"job_title",
		"base_salary",
		"department_code",
		"employment_type",
		"insurance_number",
		"tax_id",
		"hire_date",
		"is_active",
	]
	_write_header_row(ws, headers)
	for emp, person, dept in rows:
		ws.append(
			[
				emp.person_id,
				person.code,
				emp.employee_code,
				emp.job_title,
				float(emp.base_salary) if emp.base_salary is not None else None,
				dept.code if dept else None,
				emp.employment_type,
				emp.insurance_number,
				emp.tax_id,
				emp.hire_date.isoformat() if emp.hire_date else None,
				"yes" if emp.is_active else "no",
			]
		)
	buf = io.BytesIO()
	wb.save(buf)
	filename = f"payroll_employees_{business_id}_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
	return buf.getvalue(), filename


def import_employees_from_excel(
	db: Session,
	business_id: int,
	content: bytes,
	*,
	dry_run: bool = True,
	user_id: Optional[int] = None,
) -> Dict[str, Any]:
	_ensure_plugin(db, business_id)
	_validate_xlsx(content)
	wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
	ws = wb.active
	rows = list(ws.iter_rows(values_only=True))
	if not rows:
		raise ApiError("EMPTY_FILE", "فایل Excel خالی است.", http_status=400)
	headers = [_cell_str(h).lower() for h in rows[0]]
	required = {"employee_code"}
	if "person_id" not in headers and "person_code" not in headers:
		raise ApiError(
			"MISSING_COLUMNS",
			"ستون person_id یا person_code الزامی است.",
			http_status=400,
		)
	if "employee_code" not in headers:
		raise ApiError("MISSING_COLUMNS", "ستون employee_code الزامی است.", http_status=400)

	dept_by_code = {
		d.code: d.id
		for d in db.query(PayrollDepartment).filter(PayrollDepartment.business_id == business_id).all()
	}
	existing_by_code = {
		e.employee_code: e
		for e in db.query(PayrollEmployee).filter(PayrollEmployee.business_id == business_id).all()
	}
	existing_by_person = {
		e.person_id: e
		for e in db.query(PayrollEmployee).filter(PayrollEmployee.business_id == business_id).all()
	}

	created = 0
	updated = 0
	errors: List[Dict[str, Any]] = []

	for row_idx, row in enumerate(rows[1:], start=2):
		if not row or all(v is None or str(v).strip() == "" for v in row):
			continue
		data = {headers[i]: row[i] if i < len(row) else None for i in range(len(headers))}
		emp_code = _cell_str(data.get("employee_code"))
		if not emp_code:
			errors.append({"row": row_idx, "message": "employee_code خالی است"})
			continue

		person = None
		pid_raw = data.get("person_id")
		if pid_raw not in (None, ""):
			person = (
				db.query(Person)
				.filter(Person.business_id == business_id, Person.id == int(pid_raw))
				.first()
			)
		if person is None:
			pcode = _cell_str(data.get("person_code"))
			if pcode:
				person = (
					db.query(Person)
					.filter(Person.business_id == business_id, Person.code == pcode)
					.first()
				)
		if person is None:
			errors.append({"row": row_idx, "message": "شخص یافت نشد"})
			continue

		dept_id = None
		dcode = _cell_str(data.get("department_code"))
		if dcode:
			dept_id = dept_by_code.get(dcode)
			if dept_id is None:
				errors.append({"row": row_idx, "message": f"بخش {dcode} یافت نشد"})
				continue

		emp_type = _cell_str(data.get("employment_type")) or "full_time"
		if emp_type not in _EMPLOYMENT_TYPES:
			errors.append({"row": row_idx, "message": f"نوع استخدام نامعتبر: {emp_type}"})
			continue

		payload: Dict[str, Any] = {
			"person_id": person.id,
			"employee_code": emp_code,
			"job_title": _cell_str(data.get("job_title")) or None,
			"employment_type": emp_type,
			"insurance_number": _cell_str(data.get("insurance_number")) or None,
			"tax_id": _cell_str(data.get("tax_id")) or None,
		}
		bs = _cell_decimal(data.get("base_salary"))
		if bs is not None:
			payload["base_salary"] = bs
		if dept_id is not None:
			payload["department_id"] = dept_id
		hd = _cell_str(data.get("hire_date"))
		if hd:
			payload["hire_date"] = hd

		try:
			if emp_code in existing_by_code:
				if not dry_run:
					update_employee(db, business_id, existing_by_code[emp_code].id, payload, user_id=user_id)
				updated += 1
			elif person.id in existing_by_person:
				if not dry_run:
					update_employee(db, business_id, existing_by_person[person.id].id, payload, user_id=user_id)
				updated += 1
			else:
				if not dry_run:
					create_employee(db, business_id, payload, user_id=user_id)
				created += 1
		except ApiError as e:
			errors.append({"row": row_idx, "message": e.message})
		except Exception as e:
			errors.append({"row": row_idx, "message": str(e)})

	return {
		"dry_run": dry_run,
		"created": created,
		"updated": updated,
		"errors": errors,
		"error_count": len(errors),
	}


def build_run_lines_import_template(db: Session, business_id: int, run_id: int) -> Tuple[bytes, str]:
	_ensure_plugin(db, business_id)
	run = _get_run_row(db, business_id, run_id)
	items = (
		db.query(PayrollItemDefinition)
		.filter(
			PayrollItemDefinition.business_id == business_id,
			PayrollItemDefinition.is_active == True,  # noqa: E712
			PayrollItemDefinition.show_on_payslip == True,  # noqa: E712
		)
		.order_by(PayrollItemDefinition.sort_order, PayrollItemDefinition.id)
		.all()
	)
	employees = (
		db.query(PayrollEmployee, Person)
		.join(Person, Person.id == PayrollEmployee.person_id)
		.filter(
			PayrollEmployee.business_id == business_id,
			PayrollEmployee.is_active == True,  # noqa: E712
		)
		.order_by(PayrollEmployee.employee_code)
		.all()
	)
	wb = Workbook()
	ws = wb.active
	ws.title = "run_lines"
	headers = ["employee_code", "person_name"] + [f"item:{it.code}" for it in items]
	_write_header_row(ws, headers)
	for emp, person in employees:
		ws.append([emp.employee_code, person.name] + [None] * len(items))
	buf = io.BytesIO()
	wb.save(buf)
	filename = f"payroll_run_{run.code or run_id}_template.xlsx"
	return buf.getvalue(), filename


def import_run_lines_from_excel(
	db: Session,
	business_id: int,
	run_id: int,
	content: bytes,
	*,
	dry_run: bool = True,
	user_id: Optional[int] = None,
) -> Dict[str, Any]:
	_ensure_plugin(db, business_id)
	run = _get_run_row(db, business_id, run_id)
	_ensure_run_editable(run)
	settings = _get_or_create_settings(db, business_id)
	_validate_xlsx(content)

	items = (
		db.query(PayrollItemDefinition)
		.filter(
			PayrollItemDefinition.business_id == business_id,
			PayrollItemDefinition.is_active == True,  # noqa: E712
		)
		.all()
	)
	items_by_code = {it.code: it for it in items}
	employees_by_code = {
		e.employee_code: e
		for e in db.query(PayrollEmployee)
		.filter(PayrollEmployee.business_id == business_id, PayrollEmployee.is_active == True)  # noqa: E712
		.all()
	}

	wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
	ws = wb.active
	rows = list(ws.iter_rows(values_only=True))
	if not rows:
		raise ApiError("EMPTY_FILE", "فایل Excel خالی است.", http_status=400)
	headers = [_cell_str(h) for h in rows[0]]
	if "employee_code" not in headers:
		raise ApiError("MISSING_COLUMNS", "ستون employee_code الزامی است.", http_status=400)

	item_cols: Dict[int, int] = {}
	for idx, h in enumerate(headers):
		if h.startswith("item:"):
			code = h.split(":", 1)[1]
			if code in items_by_code:
				item_cols[idx] = items_by_code[code].id

	lines_payload: List[Dict[str, Any]] = []
	errors: List[Dict[str, Any]] = []

	for row_idx, row in enumerate(rows[1:], start=2):
		if not row or all(v is None or str(v).strip() == "" for v in row):
			continue
		emp_code = _cell_str(row[headers.index("employee_code")])
		emp = employees_by_code.get(emp_code)
		if not emp:
			errors.append({"row": row_idx, "message": f"پرسنل {emp_code} یافت نشد"})
			continue
		line_items: List[Dict[str, Any]] = []
		for col_idx, item_id in item_cols.items():
			if col_idx >= len(row):
				continue
			val = row[col_idx]
			if val is None or str(val).strip() == "":
				continue
			line_items.append({"item_definition_id": item_id, "amount": _cell_decimal(val)})
		lines_payload.append({"employee_id": emp.id, "items": line_items})

	if not lines_payload:
		raise ApiError("NO_DATA", "هیچ ردیف معتبری در فایل یافت نشد.", http_status=400)

	if errors and dry_run:
		return {"dry_run": True, "would_update_lines": len(lines_payload), "errors": errors, "error_count": len(errors)}

	if not dry_run:
		_persist_run_lines(db, business_id, run, lines_payload, settings)
		run.updated_at = datetime.datetime.utcnow()
		db.commit()

	return {
		"dry_run": dry_run,
		"updated_lines": len(lines_payload),
		"errors": errors,
		"error_count": len(errors),
	}
