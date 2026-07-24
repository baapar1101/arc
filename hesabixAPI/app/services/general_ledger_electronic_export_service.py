"""خروجی دفتر کل الکترونیکی مطابق قالب سازمان امور مالیاتی."""
from __future__ import annotations

import csv
import io
from typing import Any, Dict, List, Literal, Optional, Tuple

from fastapi.responses import Response
from openpyxl import Workbook
from sqlalchemy.orm import Session

from adapters.db.models.business import Business
from app.core.datetime_utils import export_filename_timestamp
from app.core.responses import ApiError
from app.services.journal_ledger_electronic_export_service import (
    CSV_ROW_THRESHOLD,
    _ensure_rial_currency,
    _resolve_export_format,
    _slugify,
    format_jalali_date_for_electronic_books,
    _amount_cell_value,
)
from app.services.journal_ledger_service import get_journal_ledger_report

# عناوین دقیق ستون‌ها مطابق قالب رسمی دفتر کل
ELECTRONIC_GENERAL_LEDGER_HEADERS: Tuple[str, ...] = (
    "کد حساب کل",
    "عنوان حساب کل",
    "کد حساب معین",
    "عنوان حساب معین",
    "کد حساب تفصیلی",
    "عنوان حساب تفصیلی",
    "گردش بدهکار (ریال)",
    "گردش بستانکار (ریال)",
    "تاریخ گردش",
)


def item_to_electronic_general_ledger_row(item: Dict[str, Any]) -> List[Any]:
    debit_value = _amount_cell_value(item.get("debit_amount"))
    credit_value = _amount_cell_value(item.get("credit_amount"))

    if debit_value is None and credit_value is None:
        return []

    return [
        item.get("general_account_code") or "",
        item.get("general_account_name") or "",
        item.get("subsidiary_account_code") or "",
        item.get("subsidiary_account_name") or "",
        item.get("detail_account_code") or "",
        item.get("detail_account_name") or "",
        debit_value,
        credit_value,
        format_jalali_date_for_electronic_books(item.get("document_date")),
    ]


def build_electronic_general_ledger_rows(items: List[Dict[str, Any]]) -> List[List[Any]]:
    rows: List[List[Any]] = []
    for item in items:
        row = item_to_electronic_general_ledger_row(item)
        if row:
            rows.append(row)
    return rows


def build_electronic_general_ledger_excel_bytes(rows: List[List[Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    try:
        ws.sheet_view.rightToLeft = True
    except Exception:
        pass

    for col_idx, header in enumerate(ELECTRONIC_GENERAL_LEDGER_HEADERS, 1):
        ws.cell(row=1, column=col_idx, value=header)

    for row_idx, row in enumerate(rows, 2):
        for col_idx, value in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if col_idx in (7, 8) and isinstance(value, int):
                cell.number_format = "0"
            if col_idx == 9 and value:
                cell.number_format = "0"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def build_electronic_general_ledger_csv_bytes(rows: List[List[Any]]) -> bytes:
    output = io.StringIO()
    writer = csv.writer(output, lineterminator="\n")
    writer.writerow(ELECTRONIC_GENERAL_LEDGER_HEADERS)
    for row in rows:
        csv_row: List[Any] = []
        for idx, value in enumerate(row):
            if idx in (6, 7):
                csv_row.append("" if value is None else int(value))
            else:
                csv_row.append("" if value is None else value)
        writer.writerow(csv_row)
    return output.getvalue().encode("utf-8-sig")


def export_electronic_general_ledger_books(
    db: Session,
    *,
    business_id: int,
    fiscal_year_id: Optional[int],
    currency_id: Optional[int],
    date_from: Optional[str],
    date_to: Optional[str],
    document_type: Optional[str],
    include_proforma: bool,
    export_format: str = "auto",
) -> Response:
    """
    خروجی دفتر کل الکترونیکی برای تمام حساب‌های کسب‌وکار در بازه انتخاب‌شده.
    فیلتر حساب (account_ids) در این خروجی نادیده گرفته می‌شود؛ مطابق اصل دفتر کل قانونی.
    """
    _ensure_rial_currency(db, currency_id)

    result = get_journal_ledger_report(
        db=db,
        business_id=business_id,
        fiscal_year_id=fiscal_year_id,
        currency_id=currency_id,
        date_from=date_from,
        date_to=date_to,
        document_type=document_type,
        include_proforma=include_proforma,
        skip=0,
        take=1,
        disable_pagination=True,
    )

    summary = result.get("summary") or {}
    if summary.get("balance_valid") is False:
        raise ApiError(
            "UNBALANCED_LEDGER",
            "دفتر کل تراز نیست؛ قبل از ارسال دفتر الکترونیکی، مغایرت‌ها را برطرف کنید.",
            http_status=400,
        )

    items = result.get("items") or []
    rows = build_electronic_general_ledger_rows(items)
    if not rows:
        raise ApiError(
            "NO_DATA",
            "در بازه انتخاب‌شده سطری برای خروجی دفتر کل الکترونیکی یافت نشد",
            http_status=404,
        )

    resolved_format = _resolve_export_format(export_format, len(rows))

    biz_name = ""
    try:
        business = db.query(Business).filter(Business.id == business_id).first()
        if business is not None:
            biz_name = business.name or ""
    except Exception:
        biz_name = ""

    base = "electronic_general_ledger"
    if biz_name:
        base += f"_{_slugify(biz_name)}"
    timestamp = export_filename_timestamp(business_id)

    if resolved_format == "csv":
        content = build_electronic_general_ledger_csv_bytes(rows)
        filename = f"{base}_{timestamp}.csv"
        media_type = "text/csv; charset=utf-8"
    else:
        content = build_electronic_general_ledger_excel_bytes(rows)
        filename = f"{base}_{timestamp}.xlsx"
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    return Response(
        content=content,
        media_type=media_type,
        headers={
            "Content-Disposition": f"attachment; filename={filename}",
            "Content-Length": str(len(content)),
            "Access-Control-Expose-Headers": "Content-Disposition",
            "X-Export-Format": resolved_format,
            "X-Export-Row-Count": str(len(rows)),
        },
    )
