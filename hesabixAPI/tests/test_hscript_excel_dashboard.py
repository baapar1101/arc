"""تست Excel و Dashboard composer برای HScript."""

from __future__ import annotations

from openpyxl import load_workbook

from app.services.hscript import run_script, validate_script
from app.services.hscript.excel_renderer import spec_to_excel_bytes


class _DummyDB:
	pass


def test_dashboard_span_and_row_break():
	src = """
report.dashboard(columns=12)
report.title("Dash")
report.kpi("A", 1, span=4)
report.kpi("B", 2, span=4)
report.kpi("C", 3, span=4)
report.row_break()
report.table([{"n": "x", "v": 1}], span=12)
"""
	assert validate_script(src)["ok"] is True
	result = run_script(src, db=_DummyDB(), business_id=1, user_id=1, preview=True)  # type: ignore[arg-type]
	assert result.ok, result.error
	assert result.spec["layout"] == "dashboard"
	assert result.spec["meta"]["dashboard_columns"] == 12
	types = [b["type"] for b in result.spec["blocks"]]
	assert "row_break" in types
	assert any(b.get("span") == 4 for b in result.spec["blocks"] if b["type"] == "kpi")


def test_excel_from_spec_has_summary_and_table_sheets():
	spec = {
		"version": 1,
		"title": "گزارش تست",
		"layout": "dashboard",
		"meta": {},
		"blocks": [
			{"type": "kpi", "label": "تعداد", "value": 12, "format": "number"},
			{
				"type": "table",
				"title": "مشتریان",
				"columns": ["name", "total"],
				"rows": [{"name": "A", "total": 10}, {"name": "B", "total": 20}],
			},
			{
				"type": "chart",
				"chart_type": "bar",
				"title": "روند",
				"data": [{"x": "A", "y": 1}, {"x": "B", "y": 2}],
			},
		],
	}
	raw = spec_to_excel_bytes(spec, business_name="Demo")
	assert raw[:2] == b"PK"  # zip/xlsx
	wb = load_workbook(filename=__import__("io").BytesIO(raw))
	names = wb.sheetnames
	assert "خلاصه" in names
	assert any("مشتریان" in n or "جدول" in n for n in names)
	assert any("نمودار" in n for n in names)
