"""تست تمپلیت و نرمال‌سازی ایمپورت اکسل کالا."""
from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.services.product_excel_import_normalize import (
    build_create_payload,
    build_update_payload,
    is_sample_import_row,
    map_headers,
    parse_bool_strict,
    parse_inventory_mode,
    provided_keys_from_raw,
    select_products_import_worksheet,
)
from app.services.product_excel_import_opening_balance import (
    OPENING_BALANCE_QUANTITY_KEY,
    WarehouseImportIndex,
    prepare_opening_balance_for_import_row,
)
from app.services.product_excel_import_spec import (
    HEADER_ALIASES,
    SAMPLE_PRODUCT_CODE,
    USER_TEMPLATE_COLUMNS,
    normalize_header,
)
from app.services.product_excel_import_template_service import (
    ProductImportLookups,
    build_products_import_workbook,
    category_full_path,
)


def test_user_template_has_category_path_not_id():
    keys = [c.key for c in USER_TEMPLATE_COLUMNS]
    assert "category_path" in keys
    assert "category_id" not in keys
    assert "default_warehouse_id" not in keys
    assert "tax_type_id" not in keys
    assert "attribute_ids" not in keys
    assert "general_barcodes" in keys
    assert "is_active" in keys
    assert "inventory_mode" in keys
    assert "price_fx_currency_code" in keys


def test_header_aliases_prefer_path_over_id():
    assert HEADER_ALIASES[normalize_header("مسیر دسته‌بندی")] == "category_path"
    assert HEADER_ALIASES[normalize_header("شناسه دسته‌بندی")] == "category_id"
    assert HEADER_ALIASES[normalize_header("بارکدهای عمومی")] == "general_barcodes"
    assert HEADER_ALIASES[normalize_header("کد ارز قیمت")] == "price_fx_currency_code"
    assert HEADER_ALIASES[normalize_header("فعال")] == "is_active"


def test_map_headers_localized():
    mapped = map_headers(["کد", "نام", "مسیر دسته‌بندی", "بارکدهای عمومی"])
    assert mapped == ["code", "name", "category_path", "general_barcodes"]


def test_parse_bool_empty_is_none():
    assert parse_bool_strict("") is None
    assert parse_bool_strict(None) is None
    assert parse_bool_strict("بله") is True
    assert parse_bool_strict("خیر") is False


def test_parse_inventory_mode():
    assert parse_inventory_mode("فله‌ای") == "bulk"
    assert parse_inventory_mode("unique") == "unique"
    assert parse_inventory_mode("") is None


def test_sample_row_detection():
    assert is_sample_import_row({"code": SAMPLE_PRODUCT_CODE, "name": "نمونه کالا"})
    assert is_sample_import_row({"code": "X", "description": "ردیف نمونه — قبل از ایمپورت"})
    assert not is_sample_import_row({"code": "P2", "name": "ماگ"})


def test_update_payload_skips_empty_cells():
    item = {
        "name": "ماگ",
        "track_inventory": True,
        "is_sales_taxable": None,
        "category_id": 4,
        "category_path": "مواد اولیه > پلاستیک",
        "_row": 2,
    }
    provided = {"name", "track_inventory", "category_path"}
    payload = build_update_payload(item, provided)
    assert payload["name"] == "ماگ"
    assert payload["track_inventory"] is True
    assert payload["category_id"] == 4
    assert "is_sales_taxable" not in payload
    assert "category_path" not in payload
    assert "_row" not in payload


def test_create_payload_drops_excel_only_and_empty():
    item = {
        "name": "ماگ",
        "category_path": "پوشاک",
        "category_id": 9,
        "track_inventory": True,
        "description": None,
        "_provided_keys": {"name"},
    }
    payload = build_create_payload(item)
    assert payload["name"] == "ماگ"
    assert payload["category_id"] == 9
    assert payload["track_inventory"] is True
    assert "category_path" not in payload
    assert "description" not in payload
    assert "_provided_keys" not in payload


def test_select_products_sheet_skips_instructions():
    wb = Workbook()
    ws_help = wb.active
    ws_help.title = "راهنما"
    ws_help["A1"] = "راهنما"
    ws_data = wb.create_sheet("کالاها")
    ws_data["A1"] = "نام"
    ws_data["B1"] = "کد"
    assert select_products_import_worksheet(wb).title == "کالاها"


def test_template_workbook_has_lookups_and_no_category_id():
    lookups = ProductImportLookups(
        category_rows=[{"path": "مواد اولیه > پلاستیک", "title_fa": "پلاستیک", "title_en": "Plastics"}],
        warehouses=[{"code": "WH1", "name": "اصلی", "is_default": True}],
        tax_types=[{"code": "VAT", "title": "ارزش افزوده"}],
        tax_units=[{"code": "U1", "name": "واحد ۱"}],
        attributes=[{"title": "رنگ", "data_type": "text"}],
        currencies=[{"code": "USD", "title": "دلار"}],
    )
    wb = build_products_import_workbook(lookups, "fa")
    names = [ws.title for ws in wb.worksheets]
    assert names[0] == "کالاها"
    assert "راهنما" in names
    assert "دسته‌بندی‌ها" in names
    assert "انبارها" in names
    headers = [c.value for c in wb["کالاها"][1]]
    assert "مسیر دسته‌بندی" in headers
    assert "شناسه دسته‌بندی" not in headers
    assert "بارکدهای عمومی" in headers
    assert "فعال" in headers
    sample_name_idx = headers.index("نام") + 1
    assert wb["کالاها"].cell(row=2, column=sample_name_idx).value == "نمونه کالا"
    cat_headers = [c.value for c in wb["دسته‌بندی‌ها"][1]]
    assert "مسیر دسته‌بندی" in cat_headers
    assert wb["دسته‌بندی‌ها"].cell(row=2, column=1).value == "مواد اولیه > پلاستیک"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    loaded = load_workbook(buf)
    assert loaded["کالاها"]["A1"].value is not None


def test_empty_opening_balance_does_not_clear_on_update():
    idx = WarehouseImportIndex([])
    existing = type("P", (), {"id": 1, "default_warehouse_id": 3, "track_inventory": True})()
    ob, errors, warnings, preview = prepare_opening_balance_for_import_row(
        item={"item_type": "کالا", "track_inventory": True},
        mapped_keys={OPENING_BALANCE_QUANTITY_KEY},
        business_id=1,
        db=None,  # type: ignore[arg-type]
        can_edit_opening_balance=True,
        is_update=True,
        existing_product=existing,
        warehouse_index=idx,
    )
    assert ob is None
    assert errors == []
    assert preview == {}


def test_provided_keys_ignore_blank():
    keys = provided_keys_from_raw({"name": "A", "code": "  ", "track_inventory": True, "_row": 2})
    assert keys == {"name", "track_inventory"}


def test_category_full_path_walks_parents():
    parent = type("C", (), {"id": 1, "parent_id": None, "title_translations": {"fa": "مواد اولیه"}})()
    child = type("C", (), {"id": 2, "parent_id": 1, "title_translations": {"fa": "پلاستیک"}})()
    by_id = {1: parent, 2: child}
    assert category_full_path(child, by_id, "fa") == "مواد اولیه > پلاستیک"
